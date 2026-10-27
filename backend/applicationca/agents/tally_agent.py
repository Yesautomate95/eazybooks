import os
import sys
import logging
import json
import time
from typing import Dict, List, Any, Optional, Literal
from datetime import datetime
from pathlib import Path
import re
import asyncio
import warnings
import requests
import base64
import mimetypes  # Add this for mime type detection
from google import genai  # New google-genai SDK
from google.genai import types as genai_types

# Pydantic and LangGraph imports
from pydantic import BaseModel, Field
from langgraph.graph import StateGraph, END

# LangChain imports
from langchain_core.messages import HumanMessage, SystemMessage
from langchain_openai import ChatOpenAI

# Document processing imports
import PyPDF2
import docx
import openpyxl
import pandas as pd

# Import Gemini key from eazybooks_secrets
try:
    from applicationca.chatbot.constants.eazybooks_secrets import gemini_client_key
    HAS_GEMINI_KEY = True
except ImportError:
    gemini_client_key = ""
    HAS_GEMINI_KEY = False
    
# Document-specific processing functions
from applicationca.agents.document_specific_functions.bank_entries import (
    extract_bank_transactions,
    match_bank_transaction_ledgers,
    generate_bank_xml
)
from applicationca.agents.document_specific_functions.sales import (
    extract_sales_transactions,
    match_sales_transaction_ledgers,
    generate_sales_xml
)
from applicationca.agents.document_specific_functions.purchase import (
    extract_purchase_transactions,
    match_purchase_transaction_ledgers,
    generate_purchase_xml
)

# Add the parent directory to the Python path
root_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.append(root_dir)

# Configure logging first
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Load environment variables
try:
    from dotenv import load_dotenv
    env_path = os.path.join(root_dir, '.env')
    logger.info(f"Looking for .env file at: {env_path}")
    load_dotenv(env_path)
    logger.info("Environment variables loaded")
except ImportError:
    logger.warning("dotenv not installed, skipping .env loading")

# API Keys and Configuration
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY", "")
if OPENAI_API_KEY:
    logger.info("OpenAI API key found in environment variables")
else:
    logger.error("OpenAI API key not found in environment variables!")

# Set Gemini API key from eazybooks_secrets if available
if HAS_GEMINI_KEY and gemini_client_key:
    GOOGLE_API_KEY = gemini_client_key
    os.environ["GOOGLE_API_KEY"] = gemini_client_key
    logger.info("Google API key found in eazybooks_secrets")
else:
    GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY", "")
    if GOOGLE_API_KEY:
        logger.info("Google API key found in environment variables")
    else:
        logger.error("Google API key not found in environment variables or eazybooks_secrets!")

# Default configurations
SELECTED_MODEL = "gpt-4.1-mini"  # Default OpenAI model
BANK_EXTRACTION_MODEL = "gpt-4.1"  # Best model for bank statement extraction (32K output tokens)
SELECTED_GEMINI_MODEL = "gemini-2.5-pro"  # Primary Google model (best for handwriting/tough images)
GEMINI_FALLBACK_MODEL = "gemini-2.5-flash"  # Fallback Google model (higher free quota)

# Extraction provider configuration
ANTHROPIC_API_KEY = os.getenv("ANTHROPIC_API_KEY", "")
EXTRACTION_PROVIDER = os.getenv("EXTRACTION_PROVIDER", "openai")
CLAUDE_EXTRACTION_MODEL = os.getenv("CLAUDE_EXTRACTION_MODEL", "claude-opus-4-7")
if ANTHROPIC_API_KEY:
    logger.info(f"Anthropic API key found, extraction provider: {EXTRACTION_PROVIDER}")

# Define the state for the agent
class AgentState(BaseModel):
    """State for the tally agent workflow."""
    # Input fields
    file_path: str
    company_name: str  # Required field without default
    user_id: Optional[str] = None  # For MappingService (Layer 1 of LedgerMatcher)

    # Processing fields
    file_type: str = ""
    extracted_text: str = ""
    document_type: str = ""  # Can be "Sales", "Purchase", or "Bank Entries"
    thought_process: List[str] = Field(default_factory=list)
    transactions: List[Dict[str, Any]] = Field(default_factory=list)
    available_ledgers: List[Dict[str, str]] = Field(default_factory=list)

    # Output fields
    tally_xml: str = ""
    final_status: str = ""
    error: str = ""
    inventory_names: List[str] = Field(default_factory=list)
    inventory_names_with_units: List[Dict[str, str]] = Field(default_factory=list)

def get_llm(temperature: float = 0.7, max_tokens: int = 1000) -> ChatOpenAI:
    """Get an OpenAI LLM instance with the specified configuration."""
    return ChatOpenAI(
        model=SELECTED_MODEL,
        openai_api_key=OPENAI_API_KEY,
        temperature=temperature,
        max_tokens=max_tokens
    )

def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from a PDF file."""
    start_time = time.time()
    logger.info(f"Starting PDF text extraction from {file_path}")
    
    text = []
    with open(file_path, 'rb') as file:
        pdf_reader = PyPDF2.PdfReader(file)
        for page in pdf_reader.pages:
            text.append(page.extract_text())
    
    elapsed_time = time.time() - start_time
    logger.info(f"PDF text extraction completed in {elapsed_time:.2f} seconds")
    return '\n'.join(text)

def extract_text_from_docx(file_path: str) -> str:
    """Extract text from a DOCX file."""
    start_time = time.time()
    logger.info(f"Starting DOCX text extraction from {file_path}")
    
    doc = docx.Document(file_path)
    text = []
    for paragraph in doc.paragraphs:
        text.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                row_text.append(cell.text.strip())
            text.append('\t'.join(row_text))
    
    elapsed_time = time.time() - start_time
    logger.info(f"DOCX text extraction completed in {elapsed_time:.2f} seconds")
    return '\n'.join(text)

def extract_text_from_xlsx(file_path: str) -> str:
    """Extract text from an XLSX file."""
    start_time = time.time()
    logger.info(f"Starting XLSX text extraction from {file_path}")
    
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    text = []
    
    for sheet in workbook.worksheets:
        text.append(f"\nSheet: {sheet.title}")
        
        # Get the maximum row and column with data
        max_row = sheet.max_row
        max_col = sheet.max_column
        
        # Extract headers
        headers = []
        for col in range(1, max_col + 1):
            cell_value = sheet.cell(row=1, column=col).value
            if cell_value is not None:
                headers.append(str(cell_value))
        if headers:
            text.append('\t'.join(headers))
        
        # Extract data rows
        for row in range(2, max_row + 1):
            row_data = []
            for col in range(1, max_col + 1):
                cell_value = sheet.cell(row=row, column=col).value
                if cell_value is not None:
                    row_data.append(str(cell_value))
            if row_data:  # Only add non-empty rows
                text.append('\t'.join(row_data))
    
    elapsed_time = time.time() - start_time
    logger.info(f"XLSX text extraction completed in {elapsed_time:.2f} seconds")
    return '\n'.join(text)

async def preprocess_file(state: AgentState) -> AgentState:
    """Preprocess the file and extract text content."""
    start_time = time.time()
    logger.info(f"Starting file preprocessing: {state.file_path}")
    print(f"\nPreprocessing file: {state.file_path}")
    
    try:
        # Get file extension
        file_extension = os.path.splitext(state.file_path)[1].lower()
        state.file_type = file_extension
        
        # Process based on file type
        if file_extension in ['.txt', '.csv', '.json']:
            # For text files, read directly
            with open(state.file_path, 'r', encoding='utf-8') as file:
                state.extracted_text = file.read()
            state.thought_process.append(f"Successfully read text file ({file_extension})")
            
        elif file_extension == '.pdf':
            # to_thread: page-by-page extraction on large PDFs is CPU-bound and
            # would otherwise block the event loop for every other request.
            state.extracted_text = await asyncio.to_thread(extract_text_from_pdf, state.file_path)
            state.thought_process.append("Successfully extracted text from PDF")
            
        elif file_extension == '.docx':
            state.extracted_text = extract_text_from_docx(state.file_path)
            state.thought_process.append("Successfully extracted text from DOCX")
            
        elif file_extension in ['.xlsx', '.xls']:
            state.extracted_text = extract_text_from_xlsx(state.file_path)
            state.thought_process.append(f"Successfully extracted text from Excel file ({file_extension})")
            
        elif file_extension in ['.jpg', '.jpeg', '.png', '.gif', '.bmp', '.tiff', '.webp']:
            # For image files, skip text extraction as Gemini will process them directly
            state.extracted_text = ""  # Empty text, Gemini will handle the image
            state.thought_process.append(f"Image file detected ({file_extension}) - will be processed by Gemini multimodal")
            print(f"Image file detected ({file_extension}) - will be processed by Gemini multimodal")
            
        else:
            state.error = f"Unsupported file type: {file_extension}"
            state.thought_process.append(f"Error: Unsupported file type {file_extension}")
            return state
        
        # Clean up the extracted text (only if we have text)
        if state.extracted_text:
            # Remove null bytes and normalize whitespace (preserve newlines for chunking)
            cleaned_text = state.extracted_text.replace('\x00', '')
            cleaned_text = re.sub(r'[^\S\n]+', ' ', cleaned_text)
            cleaned_text = re.sub(r'\n{3,}', '\n\n', cleaned_text)
            cleaned_text = cleaned_text.strip()
            state.extracted_text = cleaned_text
            
        print(f"Successfully preprocessed {file_extension} file")
            
    except Exception as e:
        state.error = f"File preprocessing failed: {str(e)}"
        state.thought_process.append(f"Error preprocessing file: {str(e)}")
        logger.error(f"Preprocessing error: {str(e)}")
    
    elapsed_time = time.time() - start_time
    logger.info(f"File preprocessing completed in {elapsed_time:.2f} seconds")
    return state

async def identify_document_type(state: AgentState) -> AgentState:
    """Identify document type using OpenAI GPT."""
    start_time = time.time()
    logger.info(f"Starting OpenAI document type identification for: {state.file_path}")
    print(f"\nIdentifying document type using OpenAI for: {state.file_path}")

    try:
        # Check if API key is valid
        if not OPENAI_API_KEY or OPENAI_API_KEY.strip() == "":
            raise ValueError("Missing OpenAI API key - please provide a valid API key")

        # Initialize OpenAI LLM
        llm = ChatOpenAI(
            model=SELECTED_MODEL,
            openai_api_key=OPENAI_API_KEY,
            temperature=0.1
        )

        # Construct the prompt with extracted text
        prompt = f"""Analyze this financial document text and classify it into one of these categories:
1. Sales Document (invoices, bills where goods/services are sold TO customers)
2. Purchase Document (bills where goods/services are bought FROM suppliers)
3. Bank Entries Document (bank statements, account statements, transaction lists)

Document text (first 3000 characters):
{state.extracted_text[:3000]}

Consider these indicators:
- Keywords like "Bank", "Account Statement", "Balance", "Debit", "Credit" suggest Bank Entries
- Keywords like "Invoice", "Bill To", "Customer", "Sales" suggest Sales
- Keywords like "Purchase Order", "Vendor", "Supplier", "Bill From" suggest Purchase
- Transaction lists with dates and amounts typically indicate Bank Entries

Respond with EXACTLY one word: "Sales", "Purchase", or "Bank Entries"."""

        # Process with OpenAI
        logger.info("Sending document to OpenAI for document type identification")
        print("Processing document with OpenAI...")

        response = await llm.ainvoke([
            SystemMessage(content="You are a document classification expert. Classify financial documents accurately based on their content."),
            HumanMessage(content=prompt)
        ])

        # Extract content from response
        content = response.content.strip().lower()

        # Map the response to document type
        if 'sales' in content:
            document_type = "Sales"
        elif 'purchase' in content:
            document_type = "Purchase"
        else:
            document_type = "Bank Entries"

        state.document_type = document_type
        state.thought_process.append(f"Identified document type using OpenAI: {document_type}")
        print(f"Document type identified: {document_type}")

    except Exception as e:
        state.error = f"Document type identification failed: {str(e)}"
        state.thought_process.append(f"Error identifying document type: {str(e)}")
        logger.error(f"Document type identification error: {str(e)}")

    elapsed_time = time.time() - start_time
    logger.info(f"Document type identification completed in {elapsed_time:.2f} seconds")
    return state

async def retrieve_ledgers(state: AgentState) -> AgentState:
    """Store the provided ledgers in the state."""
    logger.info("Storing provided ledgers")
    
    if not state.available_ledgers:
        state.error = "No ledgers provided"
        state.thought_process.append("Error: No ledgers provided")
        logger.error("No ledgers provided")
        
        # Use default ledgers as fallback
        state.available_ledgers = [
            {"name": "Cash", "parent_group": "Cash in Hand", "opening_balance": "0"},
            {"name": "Sales", "parent_group": "Sales Accounts", "opening_balance": "0"},
            {"name": "Purchase", "parent_group": "Purchase Accounts", "opening_balance": "0"}
        ]
        state.thought_process.append("Using default ledgers as fallback")
    else:
        state.thought_process.append(f"Using {len(state.available_ledgers)} provided ledgers")
        
    return state

async def extract_transactions(state: AgentState) -> AgentState:
    """Extract transactions based on document type."""
    start_time = time.time()
    logger.info(f"Starting transaction extraction for document type: {state.document_type}")
    print(f"\nExtracting transactions for {state.document_type} document...")
    
    if state.error:
        return state
    
    try:
        # Extract transactions based on document type
        if state.document_type == "Sales":
            transactions = await extract_sales_transactions(
                state.extracted_text,
                state.file_path,
                SELECTED_GEMINI_MODEL,
                GOOGLE_API_KEY,
                state.inventory_names,
                state.inventory_names_with_units,
                gemini_fallback_model=GEMINI_FALLBACK_MODEL,
                anthropic_api_key=ANTHROPIC_API_KEY,
                extraction_provider=EXTRACTION_PROVIDER,
            )
        elif state.document_type == "Purchase":
            transactions = await extract_purchase_transactions(
                state.extracted_text,
                state.file_path,
                SELECTED_GEMINI_MODEL,
                GOOGLE_API_KEY,
                state.inventory_names,
                state.inventory_names_with_units,
                anthropic_api_key=ANTHROPIC_API_KEY,
                extraction_provider=EXTRACTION_PROVIDER,
            )
        else:  # Default to Bank Entries
            transactions = await extract_bank_transactions(
                state.extracted_text,
                state.file_path,
                BANK_EXTRACTION_MODEL,
                OPENAI_API_KEY,
                anthropic_api_key=ANTHROPIC_API_KEY,
                extraction_provider=EXTRACTION_PROVIDER,
            )
        
        # Store extracted transactions
        state.transactions = transactions
        state.thought_process.append(f"Extracted {len(transactions)} transactions")
        print(f"Successfully extracted {len(transactions)} transactions")
        
    except Exception as e:
        state.error = f"Transaction extraction failed: {str(e)}"
        state.thought_process.append(f"Error extracting transactions: {str(e)}")
        logger.error(f"Transaction extraction error: {str(e)}")
    
    elapsed_time = time.time() - start_time
    logger.info(f"Transaction extraction completed in {elapsed_time:.2f} seconds")
    return state

async def match_ledgers(state: AgentState) -> AgentState:
    """Match transactions with appropriate ledgers based on document type."""
    start_time = time.time()
    logger.info(f"Starting ledger matching for document type: {state.document_type}")
    print(f"\nMatching ledgers for {state.document_type} transactions...")
    
    if state.error or not state.transactions:
        return state
    
    try:
        # Match ledgers based on document type
        if state.document_type == "Sales":
            transactions_with_ledgers = await match_sales_transaction_ledgers(
                state.transactions,
                state.available_ledgers,
                OPENAI_API_KEY,
                SELECTED_MODEL,
                user_id=state.user_id,
                company_name=state.company_name,
            )
        elif state.document_type == "Purchase":
            transactions_with_ledgers = await match_purchase_transaction_ledgers(
                state.transactions,
                state.available_ledgers,
                OPENAI_API_KEY,
                SELECTED_MODEL,
                user_id=state.user_id,
                company_name=state.company_name,
            )
        else:  # Default to Bank Entries
            transactions_with_ledgers = await match_bank_transaction_ledgers(
                state.transactions,
                state.available_ledgers,
                OPENAI_API_KEY,
                BANK_EXTRACTION_MODEL
            )
        
        # Update transactions with ledger information
        state.transactions = transactions_with_ledgers
        state.thought_process.append(f"Matched ledgers for {len(transactions_with_ledgers)} transactions")
        print(f"Successfully matched ledgers for {len(transactions_with_ledgers)} transactions")
        
    except Exception as e:
        state.error = f"Ledger matching failed: {str(e)}"
        state.thought_process.append(f"Error matching ledgers: {str(e)}")
        logger.error(f"Ledger matching error: {str(e)}")
    
    elapsed_time = time.time() - start_time
    logger.info(f"Ledger matching completed in {elapsed_time:.2f} seconds")
    return state

async def generate_xml(state: AgentState) -> AgentState:
    """Generate Tally XML based on document type and transactions."""
    start_time = time.time()
    logger.info(f"Starting XML generation for document type: {state.document_type}")
    print(f"\nGenerating Tally XML for {state.document_type}...")
    
    if state.error or not state.transactions:
        return state
    
    try:
        # Generate XML based on document type
        if state.document_type == "Sales":
            state.tally_xml = generate_sales_xml(state.transactions, state.company_name)
        elif state.document_type == "Purchase":
            state.tally_xml = generate_purchase_xml(state.transactions, state.company_name)
        else:  # Default to Bank Entries
            state.tally_xml = generate_bank_xml(state.transactions, state.company_name)
        
        if state.tally_xml:
            # Clean up the XML by removing newlines and extra whitespace
            state.tally_xml = state.tally_xml.replace('\n', '').replace('    ', '').replace('> <', '><').strip()
            state.thought_process.append("Successfully generated Tally XML")
            print("Successfully generated Tally XML")
            logger.info(f"Generated XML: {state.tally_xml[:200]}...")  # Log first 200 chars of XML
        else:
            state.error = "XML generation produced empty result"
            state.thought_process.append("Error: XML generation produced empty result")
            logger.error("XML generation produced empty result")
            
    except Exception as e:
        state.error = f"XML generation failed: {str(e)}"
        state.thought_process.append(f"Error generating XML: {str(e)}")
        logger.error(f"XML generation error: {str(e)}")
    
    elapsed_time = time.time() - start_time
    logger.info(f"XML generation completed in {elapsed_time:.2f} seconds")
    return state

def create_agent_graph():
    """Create the workflow graph for the tally agent."""
    # Create the graph
    workflow = StateGraph(AgentState)
    
    # Add the nodes
    workflow.add_node("preprocess_file", preprocess_file)
    workflow.add_node("identify_document_type", identify_document_type)
    workflow.add_node("retrieve_ledgers", retrieve_ledgers)
    workflow.add_node("extract_transactions", extract_transactions)
    workflow.add_node("match_ledgers", match_ledgers)
    workflow.add_node("generate_xml", generate_xml)

    # Define the flow between steps
    workflow.add_edge("preprocess_file", "identify_document_type")
    workflow.add_edge("identify_document_type", "retrieve_ledgers")
    workflow.add_edge("retrieve_ledgers", "extract_transactions")
    workflow.add_edge("extract_transactions", "match_ledgers")
    workflow.add_edge("match_ledgers", "generate_xml")
    workflow.add_edge("generate_xml", END)

    # Set entry point
    workflow.set_entry_point("preprocess_file")
    
    # Compile the graph
    return workflow.compile()

async def tally_agent(
    file_path: str,
    available_ledgers: Optional[List[Dict[str, Any]]] = None,
    company_name: Optional[str] = None,
    inventory_names: Optional[List[str]] = None,
    inventory_names_with_units: Optional[List[Dict[str, str]]] = None,
    user_id: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Process a document and generate Tally-compatible transactions.
    
    Args:
        file_path: Path to the document file
        available_ledgers: List of available Tally ledgers
        company_name: Name of the Tally company
        inventory_names: List of available inventory item names in Tally
        inventory_names_with_units: List of available inventory items with their units
        
    Returns:
        Dict containing processed transactions and metadata
    """
    start_time = time.time()
    logger.info(f"Starting tally agent with file: {file_path}")
    print(f"\nStarting tally agent with file: {file_path}")
    
    if not company_name:
        error_msg = "Company name is required"
        logger.error(error_msg)
        return {
            "error": error_msg,
            "file_path": file_path,
            "document_type": "Unknown",
            "final_status": "Failed",
            "transactions_count": 0,
            "thought_process": [],
            "tally_xml": None
        }
    
    # Initialize state with provided ledgers and company name
    initial_state = AgentState(
        file_path=file_path,
        available_ledgers=available_ledgers or [],
        company_name=company_name,
        inventory_names=inventory_names or [],
        inventory_names_with_units=inventory_names_with_units or [],
        user_id=user_id,
    )
    
    try:
        # Create and compile the graph
        logger.info("Creating and compiling agent graph")
        compiled_graph = create_agent_graph()
        
        # Execute the graph
        logger.info("Executing agent graph")
        final_state = await compiled_graph.ainvoke(initial_state)
        
        # Build output dictionary
        output = {
            "file_path": file_path,
            "transactions_count": 0,
            "transactions": [],  # Include raw transactions for frontend use
            "thought_process": [],
            "final_status": "Unknown",
            "tally_xml": None
        }

        # Handle different state types (dict or AgentState)
        if isinstance(final_state, dict):
            # Handle dictionary state
            output["document_type"] = final_state.get("document_type", "Unknown")
            output["transactions"] = final_state.get("transactions", [])
            output["transactions_count"] = len(output["transactions"])
            output["thought_process"] = final_state.get("thought_process", [])
            output["final_status"] = final_state.get("final_status", "Unknown")
            output["tally_xml"] = final_state.get("tally_xml")
            # Pass extracted text (truncated) for AI suggestions
            extracted = final_state.get("extracted_text", "")
            output["extracted_text"] = extracted[:3000] if extracted else ""

            # Include errors if any
            if "error" in final_state and final_state["error"]:
                output["error"] = final_state["error"]
        else:
            # Handle AgentState object
            output["document_type"] = getattr(final_state, "document_type", "Unknown")
            output["transactions"] = getattr(final_state, "transactions", [])
            output["transactions_count"] = len(output["transactions"])
            output["thought_process"] = getattr(final_state, "thought_process", [])
            output["final_status"] = getattr(final_state, "final_status", "Unknown")
            output["tally_xml"] = getattr(final_state, "tally_xml", None)
            # Pass extracted text (truncated) for AI suggestions
            extracted = getattr(final_state, "extracted_text", "")
            output["extracted_text"] = extracted[:3000] if extracted else ""

            # Include errors if any
            if hasattr(final_state, "error") and final_state.error:
                output["error"] = final_state.error
            
        elapsed_time = time.time() - start_time
        logger.info(f"Tally agent completed in {elapsed_time:.2f} seconds")
        return output
        
    except Exception as e:
        error_output = {
            "error": str(e),
            "file_path": file_path,
            "document_type": "Unknown",
            "final_status": "Failed",
            "transactions_count": 0,
            "thought_process": getattr(initial_state, "thought_process", []) if hasattr(initial_state, "thought_process") else [],
            "tally_xml": None
        }
        print(f"\nError during execution: {str(e)}")
        elapsed_time = time.time() - start_time
        logger.error(f"Tally agent failed after {elapsed_time:.2f} seconds: {str(e)}")
        return error_output

async def test_agent():
    """Test the tally agent with a sample file."""
    # Use a test file path
    file_path = r'C:\Users\Aditya\Downloads\bank_statement_sample.pdf'
    if not os.path.exists(file_path):
        file_path = input("Enter the path to a test file: ")
    
    print(f"\nTesting tally agent with:")
    print(f"File path: {file_path}\n")
    
    try:
        result = await tally_agent(file_path)
        print("\nFinal Result:")
        print(json.dumps(result, indent=2))
    except Exception as e:
        print(f"\nError during test: {str(e)}")
        print("Please ensure the file exists and is accessible.")

if __name__ == "__main__":
    asyncio.run(test_agent()) 