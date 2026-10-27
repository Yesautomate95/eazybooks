"""
Bank Reconciliation Service — orchestrates all DB operations and
the processing pipeline (parse → fetch ledger → match → store).
"""
# noqa: redeploy

import asyncio
import logging
import json
import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Dict, List, Optional, Tuple

import boto3
from botocore.exceptions import ClientError

from config.settings import settings
from applicationca.database.database_retrieval import (
    execute_query,
    get_connection,
    return_connection,
)

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# In-memory file store (used when S3 is not configured)
# ---------------------------------------------------------------------------
# Maps a "pending://<uuid>" key → (file_content_bytes, file_name)
_pending_files: Dict[str, Tuple[bytes, str]] = {}


def store_pending_file(file_content: bytes, file_name: str) -> str:
    """Store file bytes in memory and return a retrieval key."""
    import uuid
    key = f"pending://{uuid.uuid4()}"
    _pending_files[key] = (file_content, file_name)
    return key


def pop_pending_file(key: str) -> Optional[Tuple[bytes, str]]:
    """Retrieve and remove a pending file. Returns None if not found."""
    return _pending_files.pop(key, None)


# ---------------------------------------------------------------------------
# S3 helpers
# ---------------------------------------------------------------------------

def _upload_to_s3(file_content: bytes, key: str, content_type: str) -> str:
    """Upload bytes to S3 and return the URL. Raises on missing credentials."""
    if not settings.AWS_ACCESS_KEY_ID or not settings.AWS_SECRET_ACCESS_KEY:
        raise ValueError("AWS credentials not configured")
    s3 = boto3.client(
        "s3",
        aws_access_key_id=settings.AWS_ACCESS_KEY_ID,
        aws_secret_access_key=settings.AWS_SECRET_ACCESS_KEY,
        region_name=settings.AWS_DEFAULT_REGION,
    )
    bucket = settings.AWS_BUCKET_NAME if hasattr(settings, "AWS_BUCKET_NAME") else "aureus-documents"
    s3.put_object(
        Bucket=bucket,
        Key=key,
        Body=file_content,
        ContentType=content_type,
    )
    return f"https://{bucket}.s3.{settings.AWS_DEFAULT_REGION}.amazonaws.com/{key}"


# ---------------------------------------------------------------------------
# Serial number generation
# ---------------------------------------------------------------------------

def get_next_report_serial(user_id: str, today: date) -> str:
    """Return next serial string like 'BankRecon_20260405_001'."""
    date_str = today.strftime("%Y%m%d")
    prefix = f"BankRecon_{date_str}_"
    rows = execute_query(
        """
        SELECT report_name FROM bank_recon_reports
        WHERE user_id = %s AND report_name LIKE %s
        ORDER BY report_name DESC LIMIT 1
        """,
        (user_id, f"{prefix}%"),
    )
    if rows:
        last_name = rows[0]["report_name"]
        try:
            last_num = int(last_name.split("_")[-1])
        except (ValueError, IndexError):
            last_num = 0
        return f"{prefix}{str(last_num + 1).zfill(3)}"
    return f"{prefix}001"


# ---------------------------------------------------------------------------
# Report CRUD
# ---------------------------------------------------------------------------

def create_report(
    user_id: str,
    company_name: str,
    bank_ledger_name: str,
    report_name: str,
    period_from: date,
    period_to: date,
    financial_year: str,
    file_url: str,
    file_name: str,
    ledger_closing_balance_override: Optional[float] = None,
) -> Optional[Dict]:
    conn = get_connection()
    if not conn:
        return None
    try:
        cursor = conn.cursor()
        cursor.execute(
            """
            INSERT INTO bank_recon_reports
              (user_id, company_name, bank_ledger_name, report_name,
               period_from, period_to, financial_year,
               status, ledger_closing_balance_override,
               bank_statement_file_url, bank_statement_file_name)
            VALUES (%s,%s,%s,%s,%s,%s,%s,'draft',%s,%s,%s)
            RETURNING report_id, report_name, status, created_at
            """,
            (user_id, company_name, bank_ledger_name, report_name,
             period_from, period_to, financial_year,
             ledger_closing_balance_override, file_url, file_name),
        )
        row = cursor.fetchone()
        conn.commit()
        return {
            "report_id": row[0],
            "report_name": row[1],
            "status": row[2],
            "created_at": row[3].isoformat() if row[3] else None,
        }
    except Exception as e:
        conn.rollback()
        logger.error(f"create_report error: {e}")
        return None
    finally:
        return_connection(conn)


def get_reports(user_id: str, company_name: str, bank_ledger_name: str) -> List[Dict]:
    rows = execute_query(
        """
        SELECT report_id, report_name, company_name, bank_ledger_name,
               period_from, period_to, financial_year, status,
               ledger_opening_balance, bank_opening_balance,
               ledger_closing_balance, bank_closing_balance,
               reconciliation_difference, matched_count,
               unmatched_ledger_count, unmatched_bank_count,
               created_at, updated_at
        FROM bank_recon_reports
        WHERE user_id = %s AND company_name = %s AND bank_ledger_name = %s
        ORDER BY created_at DESC
        """,
        (user_id, company_name, bank_ledger_name),
    )
    return [dict(r) for r in rows] if rows else []


def get_report(report_id: int, user_id: str) -> Optional[Dict]:
    rows = execute_query(
        """
        SELECT * FROM bank_recon_reports
        WHERE report_id = %s AND user_id = %s
        """,
        (report_id, user_id),
    )
    return dict(rows[0]) if rows else None


def update_report(report_id: int, user_id: str, fields: Dict) -> bool:
    allowed = {
        "report_name", "status",
        "ledger_opening_balance", "ledger_closing_balance",
        "bank_opening_balance", "bank_closing_balance",
        "ledger_closing_balance_override",
        "reconciliation_difference",
        "ledger_credits_not_in_bank", "ledger_debits_not_in_bank",
        "bank_credits_not_in_ledger", "bank_debits_not_in_ledger",
        "total_ledger_transactions", "total_bank_transactions",
        "matched_count", "unmatched_ledger_count", "unmatched_bank_count",
    }
    update_fields = {k: v for k, v in fields.items() if k in allowed}
    if not update_fields:
        return False
    set_clause = ", ".join(f"{k} = %s" for k in update_fields)
    values = list(update_fields.values()) + [datetime.now(), report_id, user_id]
    result = execute_query(
        f"""
        UPDATE bank_recon_reports
        SET {set_clause}, updated_at = %s
        WHERE report_id = %s AND user_id = %s
        """,
        values,
    )
    return result is not None


def delete_report(report_id: int, user_id: str) -> bool:
    result = execute_query(
        "DELETE FROM bank_recon_reports WHERE report_id = %s AND user_id = %s",
        (report_id, user_id),
    )
    return result is not None


# ---------------------------------------------------------------------------
# Transaction CRUD
# ---------------------------------------------------------------------------

def bulk_insert_transactions(report_id: int, transactions: List[Dict]) -> List[int]:
    """Insert many transactions in a single batch and return their txn_ids in order."""
    if not transactions:
        return []
    conn = get_connection()
    if not conn:
        return []
    try:
        from psycopg2.extras import execute_values
        cursor = conn.cursor()
        rows = [
            (
                report_id,
                t["source"],
                t.get("date"),
                t.get("description", ""),
                t.get("reference", ""),
                t.get("debit_amount"),
                t.get("credit_amount"),
                t.get("voucher_type"),
                t.get("voucher_number"),
                t.get("tally_guid"),
                json.dumps(t.get("raw_data")) if t.get("raw_data") else None,
                t.get("is_opening_balance", False),
            )
            for t in transactions
        ]
        fetched = execute_values(
            cursor,
            """
            INSERT INTO bank_recon_transactions
              (report_id, source, transaction_date, description, reference,
               debit_amount, credit_amount, recon_status,
               voucher_type, voucher_number, tally_guid, raw_data, is_opening_balance)
            VALUES %s
            RETURNING txn_id
            """,
            rows,
            template="(%s,%s,%s,%s,%s,%s,%s,'unmatched',%s,%s,%s,%s,%s)",
            fetch=True,
        )
        ids = [row[0] for row in fetched]
        conn.commit()
        return ids
    except Exception as e:
        conn.rollback()
        logger.error(f"bulk_insert_transactions error: {e}")
        return []
    finally:
        return_connection(conn)


def get_transactions(
    report_id: int,
    source: Optional[str] = None,
    recon_status: Optional[str] = None,
) -> List[Dict]:
    base_sql = """
        SELECT txn_id, report_id, source, transaction_date, description, reference,
               debit_amount, credit_amount, recon_status, match_id,
               voucher_type, voucher_number, tally_guid, is_opening_balance, created_at,
               COALESCE(raw_data->>'manual', '') = 'true' AS is_manual,
               COALESCE(raw_data->>'adjustment', '') = 'true' AS is_adjustment
        FROM bank_recon_transactions
        WHERE report_id = %s AND is_opening_balance = FALSE
    """
    params = [report_id]
    if source:
        base_sql += " AND source = %s"
        params.append(source)
    if recon_status:
        base_sql += " AND recon_status = %s"
        params.append(recon_status)
    base_sql += " ORDER BY transaction_date, txn_id"
    rows = execute_query(base_sql, params)
    return [dict(r) for r in rows] if rows else []


def update_transaction_status(txn_id: int, recon_status: str, match_id: Optional[int] = None) -> bool:
    result = execute_query(
        "UPDATE bank_recon_transactions SET recon_status = %s, match_id = %s WHERE txn_id = %s",
        (recon_status, match_id, txn_id),
    )
    return result is not None


def mark_transactions_pending(report_id: int, txn_ids: List[int], pending: bool = True) -> int:
    """Move unmatched transactions to/from the 'pending' triage bucket.

    No Tally writes — this is pure recon triage ("set this row aside for later").
    A pending row drops out of the unmatched lists and the reconciliation
    difference. Only `unmatched` <-> `pending` transitions are allowed, so matched
    rows are never disturbed. Returns the number of rows changed.
    """
    if not txn_ids:
        return 0
    target = "pending" if pending else "unmatched"
    source_status = "unmatched" if pending else "pending"
    conn = get_connection()
    if not conn:
        return 0
    try:
        cursor = conn.cursor()
        placeholders = ",".join(["%s"] * len(txn_ids))
        cursor.execute(
            f"""
            UPDATE bank_recon_transactions
            SET recon_status = %s
            WHERE report_id = %s AND recon_status = %s AND txn_id IN ({placeholders})
            """,
            [target, report_id, source_status, *txn_ids],
        )
        n = cursor.rowcount
        conn.commit()
        return n
    except Exception as e:
        conn.rollback()
        logger.error(f"mark_transactions_pending error: {e}")
        return 0
    finally:
        return_connection(conn)


# ---------------------------------------------------------------------------
# Match CRUD (group-based: a match may hold N ledger + M bank members)
# ---------------------------------------------------------------------------

def bulk_create_match_groups(
    report_id: int,
    groups: List[Dict],
    ledger_ids: List[int],
    bank_ids: List[int],
) -> Dict[str, int]:
    """Persist the groups returned by run_group_matching.

    Each group is {book_idxs, bank_idxs, confidence, tier, reason}. The idxs point
    into the inserted-row order, so ledger_ids[idx] / bank_ids[idx] recover the
    persisted txn_ids. (run_group_matching names the ledger side "book_idxs" — it is
    ERP-agnostic; here that side is the Tally ledger.) A group becomes one match
    header + N member rows. Tier drives the initial review state:
        tier 'auto'    -> status 'accepted'       (auto-reconciled, badge A)
        tier 'suggest' -> status 'pending_review' (needs review, badge S)
    All member txns are flipped to 'auto_matched' (tentative); rejecting a match
    returns them to 'unmatched'. Returns {total, auto, suggest}.
    """
    if not groups:
        return {"total": 0, "auto": 0, "suggest": 0}
    conn = get_connection()
    if not conn:
        return {"total": 0, "auto": 0, "suggest": 0}
    try:
        from psycopg2.extras import execute_values
        cursor = conn.cursor()
        auto_n = 0
        for g in groups:
            ledger_txn_ids = [ledger_ids[i] for i in g["book_idxs"]]
            bank_txn_ids = [bank_ids[j] for j in g["bank_idxs"]]
            tier = g.get("tier", "suggest")
            status = "accepted" if tier == "auto" else "pending_review"
            if tier == "auto":
                auto_n += 1
            # denormalized 1:1 columns kept only for the trivial single+single case
            single_ledger = ledger_txn_ids[0] if len(ledger_txn_ids) == 1 else None
            single_bank = bank_txn_ids[0] if len(bank_txn_ids) == 1 else None
            cursor.execute(
                """
                INSERT INTO bank_recon_matches
                  (report_id, match_type, status, tier, ledger_txn_id, bank_txn_id,
                   confidence_score, match_reason, notes)
                VALUES (%s,'automatic',%s,%s,%s,%s,%s,%s,'')
                RETURNING match_id
                """,
                (report_id, status, tier, single_ledger, single_bank,
                 g.get("confidence"), g.get("reason", "")),
            )
            match_id = cursor.fetchone()[0]
            member_rows = (
                [(match_id, t, "ledger") for t in ledger_txn_ids]
                + [(match_id, t, "bank") for t in bank_txn_ids]
            )
            execute_values(
                cursor,
                "INSERT INTO bank_recon_match_members (match_id, txn_id, side) VALUES %s",
                member_rows,
            )
            cursor.executemany(
                "UPDATE bank_recon_transactions SET recon_status='auto_matched', match_id=%s WHERE txn_id=%s",
                [(match_id, t) for t in (ledger_txn_ids + bank_txn_ids)],
            )
        conn.commit()
        return {"total": len(groups), "auto": auto_n, "suggest": len(groups) - auto_n}
    except Exception as e:
        conn.rollback()
        logger.error(f"bulk_create_match_groups error: {e}")
        return {"total": 0, "auto": 0, "suggest": 0}
    finally:
        return_connection(conn)


def get_matches(report_id: int, status: Optional[str] = None) -> List[Dict]:
    """Return match groups for a report. Each match carries member arrays
    (`bank_transactions` / `ledger_transactions`) so 1-to-many and many-to-many
    matches render correctly. Singular `bank_transaction` / `ledger_transaction`
    (the first member) are also included for back-compat with 1:1 consumers.
    """
    sql = """
        SELECT match_id, match_type, tier, status, confidence_score,
               match_reason, notes, created_at, reviewed_at
        FROM bank_recon_matches
        WHERE report_id = %s
    """
    params = [report_id]
    if status:
        sql += " AND status = %s"
        params.append(status)
    sql += " ORDER BY created_at"
    headers = execute_query(sql, params)
    if not headers:
        return []

    match_ids = [h["match_id"] for h in headers]
    placeholders = ",".join(["%s"] * len(match_ids))
    members = execute_query(
        f"""
        SELECT mm.match_id, mm.side,
               t.txn_id, t.transaction_date, t.description, t.reference,
               t.debit_amount, t.credit_amount, t.voucher_type, t.voucher_number
        FROM bank_recon_match_members mm
        JOIN bank_recon_transactions t ON t.txn_id = mm.txn_id
        WHERE mm.match_id IN ({placeholders})
        ORDER BY t.transaction_date, t.txn_id
        """,
        match_ids,
    )

    from collections import defaultdict
    grouped: Dict[int, Dict[str, List]] = defaultdict(lambda: {"bank": [], "ledger": []})
    for m in members or []:
        txn = {
            "txn_id": m["txn_id"],
            "transaction_date": m["transaction_date"].isoformat() if m["transaction_date"] else None,
            "description": m["description"],
            "reference": m["reference"],
            "debit_amount": float(m["debit_amount"]) if m["debit_amount"] else None,
            "credit_amount": float(m["credit_amount"]) if m["credit_amount"] else None,
        }
        if m["side"] == "ledger":
            txn["voucher_type"] = m["voucher_type"]
            txn["voucher_number"] = m["voucher_number"]
        grouped[m["match_id"]][m["side"]].append(txn)

    result = []
    for h in headers:
        g = grouped.get(h["match_id"], {"bank": [], "ledger": []})
        result.append({
            "match_id": h["match_id"],
            "match_type": h["match_type"],
            "tier": h.get("tier"),
            "status": h["status"],
            "confidence_score": float(h["confidence_score"]) if h["confidence_score"] else None,
            "match_reason": h["match_reason"],
            "notes": h["notes"],
            "created_at": h["created_at"].isoformat() if h["created_at"] else None,
            "reviewed_at": h["reviewed_at"].isoformat() if h["reviewed_at"] else None,
            "bank_transactions": g["bank"],
            "ledger_transactions": g["ledger"],
            "bank_transaction": g["bank"][0] if g["bank"] else None,
            "ledger_transaction": g["ledger"][0] if g["ledger"] else None,
        })
    return result


def update_match_status(match_id: int, status: str) -> Optional[Dict]:
    """Accept or reject a match group. Flips ALL member txns; on reject they
    return to unmatched."""
    conn = get_connection()
    if not conn:
        logger.error("update_match_status: could not get DB connection")
        return None
    try:
        cursor = conn.cursor()
        cursor.execute(
            """
            UPDATE bank_recon_matches
            SET status = %s, reviewed_at = NOW()
            WHERE match_id = %s
            RETURNING match_type
            """,
            (status, match_id),
        )
        row = cursor.fetchone()
        if not row:
            conn.rollback()
            logger.warning(f"update_match_status: match_id={match_id} not found")
            return None
        match_type = row[0]

        cursor.execute(
            "SELECT txn_id FROM bank_recon_match_members WHERE match_id = %s",
            (match_id,),
        )
        member_ids = [r[0] for r in cursor.fetchall()]

        if member_ids:
            if status == "accepted":
                new_status = "manual_matched" if match_type == "manual" else "auto_matched"
                cursor.executemany(
                    "UPDATE bank_recon_transactions SET recon_status=%s, match_id=%s WHERE txn_id=%s",
                    [(new_status, match_id, t) for t in member_ids],
                )
            elif status == "rejected":
                cursor.executemany(
                    "UPDATE bank_recon_transactions SET recon_status='unmatched', match_id=NULL WHERE txn_id=%s",
                    [(t,) for t in member_ids],
                )

        conn.commit()
        return {"match_id": match_id, "status": status}
    except Exception as e:
        conn.rollback()
        logger.error(f"update_match_status error (match_id={match_id}, status={status}): {e}")
        return None
    finally:
        return_connection(conn)


def bulk_update_match_status(match_ids: List[int], status: str) -> int:
    """Bulk accept/reject. Returns count of updated matches."""
    updated = 0
    for mid in match_ids:
        result = update_match_status(mid, status)
        if result:
            updated += 1
    return updated


def create_manual_match_group(
    report_id: int,
    ledger_txn_ids: List[int],
    bank_txn_ids: List[int],
    notes: str = "",
) -> Optional[Dict]:
    """User-driven many-to-many match. Creates an accepted manual match group and
    flips all member txns to manual_matched. Either side may be empty (e.g. an
    all-ledger contra), but not both."""
    ledger_txn_ids = list(ledger_txn_ids or [])
    bank_txn_ids = list(bank_txn_ids or [])
    if not ledger_txn_ids and not bank_txn_ids:
        return None
    conn = get_connection()
    if not conn:
        return None
    try:
        from psycopg2.extras import execute_values
        cursor = conn.cursor()
        single_ledger = ledger_txn_ids[0] if len(ledger_txn_ids) == 1 else None
        single_bank = bank_txn_ids[0] if len(bank_txn_ids) == 1 else None
        cursor.execute(
            """
            INSERT INTO bank_recon_matches
              (report_id, match_type, status, tier, ledger_txn_id, bank_txn_id,
               confidence_score, match_reason, notes, reviewed_at)
            VALUES (%s,'manual','accepted','manual',%s,%s,1.0,%s,%s,NOW())
            RETURNING match_id
            """,
            (report_id, single_ledger, single_bank, "Manual match by user", notes),
        )
        match_id = cursor.fetchone()[0]
        member_rows = (
            [(match_id, t, "ledger") for t in ledger_txn_ids]
            + [(match_id, t, "bank") for t in bank_txn_ids]
        )
        execute_values(
            cursor,
            "INSERT INTO bank_recon_match_members (match_id, txn_id, side) VALUES %s",
            member_rows,
        )
        cursor.executemany(
            "UPDATE bank_recon_transactions SET recon_status='manual_matched', match_id=%s WHERE txn_id=%s",
            [(match_id, t) for t in (ledger_txn_ids + bank_txn_ids)],
        )
        conn.commit()
        return {"match_id": match_id, "status": "accepted"}
    except Exception as e:
        conn.rollback()
        logger.error(f"create_manual_match_group error: {e}")
        return None
    finally:
        return_connection(conn)


def create_manual_match(
    report_id: int,
    ledger_txn_id: int,
    bank_txn_id: int,
    notes: str = "",
) -> Optional[Dict]:
    """Back-compat 1:1 wrapper over create_manual_match_group."""
    return create_manual_match_group(report_id, [ledger_txn_id], [bank_txn_id], notes)


def delete_match(match_id: int) -> bool:
    """Remove a match group and return all member transactions to unmatched."""
    conn = get_connection()
    if not conn:
        return False
    try:
        cursor = conn.cursor()
        cursor.execute(
            "SELECT txn_id FROM bank_recon_match_members WHERE match_id = %s",
            (match_id,),
        )
        member_ids = [r[0] for r in cursor.fetchall()]
        cursor.execute("DELETE FROM bank_recon_matches WHERE match_id = %s", (match_id,))
        if cursor.rowcount == 0:
            conn.rollback()
            return False
        if member_ids:
            cursor.executemany(
                "UPDATE bank_recon_transactions SET recon_status='unmatched', match_id=NULL WHERE txn_id=%s",
                [(t,) for t in member_ids],
            )
        conn.commit()
        return True
    except Exception as e:
        conn.rollback()
        logger.error(f"delete_match error: {e}")
        return False
    finally:
        return_connection(conn)


# ---------------------------------------------------------------------------
# Adjusting entries (create a missing ledger entry + post the voucher to Tally)
# ---------------------------------------------------------------------------

# Default Tally parent group for an auto-created counter ledger, by direction.
_RECEIPT_COUNTER_GROUP = "Indirect Incomes"   # money in  -> credit an income/other
_PAYMENT_COUNTER_GROUP = "Indirect Expenses"  # money out -> debit an expense/other


def _coerce_date(value) -> Optional[date]:
    """Coerce a parsed/posted date (date/datetime/str) to a date object."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    from applicationca.agents.document_specific_functions.bank_reconciliation import _parse_date
    return _parse_date(value)


async def create_adjustment(
    report_id: int,
    user_id: str,
    firebase_uid: str,
    *,
    date_str: str,
    description: str,
    voucher_type: str,
    amount: float,
    party_ledger: str,
    party_ledger_group: Optional[str] = None,
    reference: str = "",
    bank_txn_ids: Optional[List[int]] = None,
    auto_create_masters: bool = True,
    auto_match: bool = False,
) -> Dict:
    """Create a missing ledger entry during manual reconciliation, POST the voucher
    to Tally through the desktop bridge, and record it as a ledger transaction.

    Mirrors the QB adjustment flow; the only ERP-specific swap is the write path:
    build a Tally voucher (Receipt/Payment/Contra/Journal) and push it over the
    WebSocket bridge instead of calling the QBO API. The counter ledger is
    auto-created in Tally if it is missing (auto_create_masters).

    Accounting note (same invariant as QB): the effective ledger closing balance is
    shifted by the signed effect of the posted entry (Receipt=+, else -). Posting an
    adjustment corrects the books; it does NOT move the true book-vs-bank gap, so
    this shift plus the new (matched or unmatched) row leaves the reconciliation
    difference invariant either way.
    """
    from applicationca.contoller.websocket_controller import manager
    from applicationca.services.tally_connector import tally_connector

    report = get_report(report_id, user_id)
    if not report:
        return {"error": "Report not found"}

    voucher_type = (voucher_type or "").strip().title()
    if voucher_type not in ("Receipt", "Payment", "Contra", "Journal"):
        return {"error": "voucher_type must be Receipt, Payment, Contra or Journal"}

    party_ledger = (party_ledger or "").strip()
    if not party_ledger:
        return {"error": "A counter ledger (party_ledger) is required"}

    try:
        amt = round(abs(float(amount)), 2)
    except (TypeError, ValueError):
        return {"error": "Invalid amount"}
    if amt == 0:
        return {"error": "Amount must be non-zero"}

    company = report["company_name"]
    bank_ledger = report["bank_ledger_name"]
    is_inflow = voucher_type == "Receipt"
    created_masters: List[str] = []

    # ── Step 1: ensure the counter ledger exists in Tally (production-safe path) ──
    if auto_create_masters:
        default_group = _RECEIPT_COUNTER_GROUP if is_inflow else _PAYMENT_COUNTER_GROUP
        group = (party_ledger_group or "").strip() or default_group
        try:
            existing = await manager.get_tally_ledgers(firebase_uid, company)
            existing_names = {
                ((l.get("name") if isinstance(l, dict) else l) or "").strip().lower()
                for l in (existing or [])
            }
        except Exception as e:
            logger.warning(f"create_adjustment: could not fetch ledgers to dedupe ({e})")
            existing_names = set()
        if party_ledger.lower() not in existing_names:
            try:
                await manager.create_tally_ledger(firebase_uid, {
                    "name": party_ledger,
                    "parent_group": group,
                    "company": company,
                    "opening_balance": "0",
                })
                created_masters.append(party_ledger)
            except Exception as e:
                detail = getattr(e, "detail", None) or str(e)
                logger.error(f"create_adjustment: failed to create counter ledger '{party_ledger}': {detail}")
                return {"error": f"Could not create counter ledger '{party_ledger}' in Tally: {detail}"}

    # ── Step 2: build the voucher XML and push it to Tally ──
    voucher = {
        "date": date_str,
        "narration": description or "",
        "amount": amt,
        "voucher_type": voucher_type,
        "party_ledger": party_ledger,
        "bank_ledger": bank_ledger,
        "reference": reference or "",
    }
    xml = tally_connector.build_bank_vouchers_xml(company, bank_ledger, [voucher])
    if not xml:
        return {"error": "Could not build a valid voucher (check date and amount)"}

    try:
        push = await manager.push_to_tally(firebase_uid, company, xml)
    except Exception as e:
        detail = getattr(e, "detail", None) or str(e)
        return {"error": f"Failed to push adjustment to Tally: {detail}"}

    # The message was delivered — but did Tally actually create the voucher?
    if push.get("error"):
        return {"error": push["error"], "push": push}
    exceptions = (push.get("exception") or push.get("errors") or 0)
    if exceptions and exceptions > 0:
        return {"error": f"Tally rejected the voucher ({exceptions} exception(s))", "push": push}
    created = push.get("created", 0) or 0
    altered = push.get("altered", 0) or 0
    if not (created or altered or push.get("success")):
        return {"error": "Tally did not create the voucher", "push": push}

    # The created-voucher id field name has varied across the bridge; accept any.
    tally_vch_id = (
        push.get("lastvchid")
        or push.get("lastVoucherNumber")
        or push.get("lastVchId")
        or ""
    )

    # ── Step 3: record it as a ledger transaction (Receipt = inflow = debit) ──
    ledger_insert = {
        "source": "ledger",
        "date": _coerce_date(date_str),
        "description": description or "",
        "reference": reference or tally_vch_id or "",
        "debit_amount": amt if is_inflow else None,
        "credit_amount": None if is_inflow else amt,
        "voucher_type": voucher_type,
        "voucher_number": tally_vch_id or "",
        "tally_guid": "",
        "raw_data": {
            "adjustment": "true",
            "voucher_type": voucher_type,
            "party_ledger": party_ledger,
        },
    }
    new_ids = bulk_insert_transactions(report_id, [ledger_insert])
    new_ledger_id = new_ids[0] if new_ids else None

    # ── Step 4: keep the reconciliation difference invariant: shift effective CB ──
    delta = amt if is_inflow else -amt
    if report.get("ledger_closing_balance_override") is not None:
        update_report(report_id, user_id, {
            "ledger_closing_balance_override": float(report["ledger_closing_balance_override"]) + delta,
        })
    else:
        update_report(report_id, user_id, {
            "ledger_closing_balance": float(report.get("ledger_closing_balance") or 0) + delta,
        })

    # ── Step 5: optional immediate match to the bank line(s) it explains ──
    match = None
    if auto_match and new_ledger_id and bank_txn_ids:
        match = create_manual_match_group(
            report_id, [new_ledger_id], list(bank_txn_ids), notes="Adjustment"
        )

    return {
        "ok": True,
        "tally_voucher_id": tally_vch_id,
        "voucher_type": voucher_type,
        "ledger_txn_id": new_ledger_id,
        "match": match,
        "created_masters": created_masters,
    }


# ---------------------------------------------------------------------------
# Manual bank-statement entries (statement rows the parser missed)
# ---------------------------------------------------------------------------

def _validate_bank_entry_amounts(debit_amount, credit_amount):
    """Exactly one of debit/credit must be a positive amount. Returns
    (debit, credit, error)."""
    try:
        debit = round(abs(float(debit_amount)), 2) if debit_amount is not None else None
        credit = round(abs(float(credit_amount)), 2) if credit_amount is not None else None
    except (TypeError, ValueError):
        return None, None, "Invalid amount"
    debit = debit or None
    credit = credit or None
    if bool(debit) == bool(credit):
        return None, None, "Provide exactly one of debit_amount or credit_amount (> 0)"
    return debit, credit, None


def create_manual_bank_entry(
    report_id: int,
    *,
    date_str: str,
    description: str,
    reference: str = "",
    debit_amount: Optional[float] = None,
    credit_amount: Optional[float] = None,
) -> Dict:
    """Add a bank-statement row the parser missed. Local-only — the bank can't be
    written to, so the row just joins the statement side of the recon (flagged
    `manual` so the UI can badge and edit/delete it).

    No balance shift: the statement closing balance already reflects every real
    bank movement, including ones the parser dropped, so the new unmatched row
    slots into the bank-side summary categories and the difference self-corrects.
    """
    tdate = _coerce_date(date_str)
    if not tdate:
        return {"error": "Invalid date"}
    debit, credit, err = _validate_bank_entry_amounts(debit_amount, credit_amount)
    if err:
        return {"error": err}

    new_ids = bulk_insert_transactions(report_id, [{
        "source": "bank_statement",
        "date": tdate,
        "description": description or "",
        "reference": reference or "",
        "debit_amount": debit,
        "credit_amount": credit,
        "raw_data": {"manual": "true"},
    }])
    if not new_ids:
        return {"error": "Failed to add bank entry"}
    return {"ok": True, "bank_txn_id": new_ids[0]}


def _get_editable_manual_bank_entry(cursor, report_id: int, txn_id: int) -> Optional[str]:
    """Lock the row and verify it is a manual, still-unreconciled bank entry.
    Returns an error string, or None when the row is editable."""
    cursor.execute(
        """
        SELECT source, recon_status, COALESCE(raw_data->>'manual', '') = 'true'
        FROM bank_recon_transactions
        WHERE report_id = %s AND txn_id = %s
        FOR UPDATE
        """,
        (report_id, txn_id),
    )
    row = cursor.fetchone()
    if not row:
        return "Bank entry not found"
    source, recon_status, is_manual = row
    if source != "bank_statement" or not is_manual:
        return "Only manually added bank entries can be modified"
    if recon_status not in ("unmatched", "pending"):
        return "Entry is part of a match — remove the match first"
    return None


def update_manual_bank_entry(
    report_id: int,
    txn_id: int,
    *,
    date_str: str,
    description: str,
    reference: str = "",
    debit_amount: Optional[float] = None,
    credit_amount: Optional[float] = None,
) -> Dict:
    """Full-replace edit of a manual bank entry (only while unmatched/pending)."""
    tdate = _coerce_date(date_str)
    if not tdate:
        return {"error": "Invalid date"}
    debit, credit, err = _validate_bank_entry_amounts(debit_amount, credit_amount)
    if err:
        return {"error": err}

    conn = get_connection()
    if not conn:
        return {"error": "Database unavailable"}
    try:
        cursor = conn.cursor()
        err = _get_editable_manual_bank_entry(cursor, report_id, txn_id)
        if err:
            conn.rollback()
            return {"error": err}
        cursor.execute(
            """
            UPDATE bank_recon_transactions
            SET transaction_date = %s, description = %s, reference = %s,
                debit_amount = %s, credit_amount = %s
            WHERE report_id = %s AND txn_id = %s
            """,
            (tdate, description or "", reference or "", debit, credit, report_id, txn_id),
        )
        conn.commit()
        return {"ok": True, "bank_txn_id": txn_id}
    except Exception as e:
        conn.rollback()
        logger.error(f"update_manual_bank_entry error: {e}")
        return {"error": "Failed to update bank entry"}
    finally:
        return_connection(conn)


def delete_manual_bank_entry(report_id: int, txn_id: int) -> Dict:
    """Delete a manual bank entry (only while unmatched/pending)."""
    conn = get_connection()
    if not conn:
        return {"error": "Database unavailable"}
    try:
        cursor = conn.cursor()
        err = _get_editable_manual_bank_entry(cursor, report_id, txn_id)
        if err:
            conn.rollback()
            return {"error": err}
        cursor.execute(
            "DELETE FROM bank_recon_transactions WHERE report_id = %s AND txn_id = %s",
            (report_id, txn_id),
        )
        conn.commit()
        return {"ok": True}
    except Exception as e:
        conn.rollback()
        logger.error(f"delete_manual_bank_entry error: {e}")
        return {"error": "Failed to delete bank entry"}
    finally:
        return_connection(conn)


# ---------------------------------------------------------------------------
# Reconciliation Summary
# ---------------------------------------------------------------------------

def _effective_ledger_cb(report: Dict) -> Decimal:
    """User override wins over the auto-derived ledger closing balance."""
    override = report.get("ledger_closing_balance_override")
    if override is not None:
        return Decimal(str(override))
    return Decimal(str(report.get("ledger_closing_balance") or 0))


def get_summary(report_id: int, user_id: str) -> Optional[Dict]:
    from applicationca.agents.document_specific_functions.bank_reconciliation import (
        compute_reconciliation_summary,
    )

    report = get_report(report_id, user_id)
    if not report:
        return None

    # Unmatched ledger + bank transactions (pending rows are set aside, excluded)
    unmatched_ledger = get_transactions(report_id, source="ledger", recon_status="unmatched")
    unmatched_bank = get_transactions(report_id, source="bank_statement", recon_status="unmatched")

    lcb = _effective_ledger_cb(report)
    bcb = Decimal(str(report.get("bank_closing_balance") or 0))

    summary = compute_reconciliation_summary(lcb, bcb, unmatched_ledger, unmatched_bank)

    pending_count = execute_query(
        "SELECT COUNT(*) as cnt FROM bank_recon_matches WHERE report_id = %s AND status = 'pending_review'",
        (report_id,),
    )
    accepted_count = execute_query(
        "SELECT COUNT(*) as cnt FROM bank_recon_matches WHERE report_id = %s AND status = 'accepted'",
        (report_id,),
    )
    pending_txn_count = execute_query(
        "SELECT COUNT(*) as cnt FROM bank_recon_transactions WHERE report_id = %s AND recon_status = 'pending'",
        (report_id,),
    )

    # Automatic reconciliation % = share of (non-opening) transactions reconciled by
    # the engine at the 'auto' tier — the headline metric.
    total_txn = execute_query(
        "SELECT COUNT(*) as cnt FROM bank_recon_transactions WHERE report_id = %s AND is_opening_balance = FALSE",
        (report_id,),
    )
    auto_txn = execute_query(
        """
        SELECT COUNT(DISTINCT mm.txn_id) as cnt
        FROM bank_recon_match_members mm
        JOIN bank_recon_matches m ON m.match_id = mm.match_id
        WHERE m.report_id = %s AND m.tier = 'auto'
        """,
        (report_id,),
    )
    total_n = total_txn[0]["cnt"] if total_txn else 0
    auto_n = auto_txn[0]["cnt"] if auto_txn else 0
    automatic_pct = round((auto_n / total_n) * 100, 2) if total_n else 0.0

    return {
        "report_id": report_id,
        "report_name": report.get("report_name"),
        "status": report.get("status"),
        "company_name": report.get("company_name"),
        "bank_ledger_name": report.get("bank_ledger_name"),
        "period_from": report["period_from"].isoformat() if report.get("period_from") else None,
        "period_to": report["period_to"].isoformat() if report.get("period_to") else None,
        "financial_year": report.get("financial_year"),
        "ledger_opening_balance": float(report.get("ledger_opening_balance") or 0),
        "ledger_closing_balance": float(lcb),
        "bank_opening_balance": float(report.get("bank_opening_balance") or 0),
        "bank_closing_balance": float(bcb),
        "summary": summary,
        "counts": {
            "total_matched": report.get("matched_count", 0),
            "unmatched_ledger": len(unmatched_ledger),
            "unmatched_bank": len(unmatched_bank),
            "pending_review": pending_count[0]["cnt"] if pending_count else 0,
            "accepted": accepted_count[0]["cnt"] if accepted_count else 0,
            "pending_transactions": pending_txn_count[0]["cnt"] if pending_txn_count else 0,
            "automatic_pct": automatic_pct,
        },
    }


# ---------------------------------------------------------------------------
# Processing Pipeline
# ---------------------------------------------------------------------------

async def process_report(report_id: int, user_id: str, firebase_uid: str) -> Dict:
    """
    Full processing pipeline:
      1. Parse bank statement from S3
      2. Fetch ledger transactions from Tally via WebSocket
      3. Insert both into bank_recon_transactions
      4. Run automatic matching
      5. Insert matches
      6. Update report counts and balances
    """
    from applicationca.contoller.websocket_controller import manager
    from applicationca.agents.document_specific_functions.bank_reconciliation import (
        parse_bank_statement_with_llm,
        run_group_matching,
    )

    report = get_report(report_id, user_id)
    if not report:
        return {"error": "Report not found"}

    # Idempotency: if already processing or done, return existing state.
    # Prevents double-firing (frontend timeout + retry) from re-popping a
    # memory-only file and blowing up with "file not found" on the 2nd call.
    # awaiting_tally is NOT terminal — it means "ready to retry now that
    # Tally is online" — so it falls through into the normal pipeline.
    current_status = report.get("status")
    if current_status in ("processing", "processed", "completed"):
        existing_matches = get_matches(report_id) if "get_matches" in globals() else []
        return {
            "status": current_status,
            "message": f"Report already {current_status}",
            "matches_count": len(existing_matches) if existing_matches else None,
            "idempotent": True,
        }

    update_report(report_id, user_id, {"status": "processing"})

    try:
        # ── Step 1: Get bank statement file ──
        file_url = report.get("bank_statement_file_url", "")
        file_name = report.get("bank_statement_file_name", "statement.pdf")
        period_from = report["period_from"]
        period_to = report["period_to"]

        file_content: Optional[bytes] = None

        if file_url and file_url.startswith("pending://"):
            pending = pop_pending_file(file_url)
            if pending:
                file_content, file_name = pending
            else:
                update_report(report_id, user_id, {"status": "error"})
                return {"error": "Uploaded file not found in memory. Please re-upload the statement."}
        elif file_url:
            try:
                s3 = boto3.client(
                    "s3",
                    aws_access_key_id=settings.AWS_ACCESS_KEY_ID,
                    aws_secret_access_key=settings.AWS_SECRET_ACCESS_KEY,
                    region_name=settings.AWS_DEFAULT_REGION,
                )
                bucket = settings.AWS_BUCKET_NAME if hasattr(settings, "AWS_BUCKET_NAME") else "aureus-documents"
                s3_key = file_url.split(f"{bucket}.s3.{settings.AWS_DEFAULT_REGION}.amazonaws.com/")[-1]
                obj = s3.get_object(Bucket=bucket, Key=s3_key)
                file_content = obj["Body"].read()
            except Exception as e:
                logger.error(f"S3 download failed: {e}")
                update_report(report_id, user_id, {"status": "error"})
                return {"error": "Could not retrieve bank statement file from storage."}
        else:
            update_report(report_id, user_id, {"status": "error"})
            return {"error": "No bank statement file associated with this report."}

        # ── Steps 1 & 2 in parallel: parse bank statement + fetch Tally ledger ──
        company = report["company_name"]
        bank_ledger = report["bank_ledger_name"]
        logger.info(
            f"Parallel: parsing '{file_name}' + fetching Tally ledger '{bank_ledger}' "
            f"[{period_from.strftime('%Y%m%d')} – {period_to.strftime('%Y%m%d')}]"
        )
        parsed, ledger_data = await asyncio.gather(
            parse_bank_statement_with_llm(
                file_content, file_name,
                period_from=period_from,
                period_to=period_to,
            ),
            manager.get_tally_ledger_transactions(
                firebase_uid,
                company=company,
                ledger_name=bank_ledger,
                from_date=period_from.strftime("%Y%m%d"),
                to_date=period_to.strftime("%Y%m%d"),
            ),
        )
        bank_txns_raw = parsed["transactions"]
        bank_ob = parsed.get("opening_balance")
        bank_cb = parsed.get("closing_balance")
        logger.info(f"Parsed {len(bank_txns_raw)} bank transactions (OB={bank_ob}, CB={bank_cb})")

        ledger_ob = None
        ledger_cb = None
        ledger_txns_raw = []
        if isinstance(ledger_data, dict):
            ledger_ob = Decimal(str(ledger_data.get("opening_balance") or 0)) or None
            ledger_cb = Decimal(str(ledger_data.get("closing_balance") or 0)) or None
            ledger_txns_raw = ledger_data.get("transactions", [])
        elif isinstance(ledger_data, list):
            ledger_txns_raw = ledger_data

        # ── Step 3: Insert bank statement transactions ──
        bank_inserts = []
        for t in bank_txns_raw:
            bank_inserts.append({
                "source": "bank_statement",
                "date": t.get("date"),
                "description": t.get("description", ""),
                "reference": t.get("reference", ""),
                "debit_amount": float(t["debit_amount"]) if t.get("debit_amount") else None,
                "credit_amount": float(t["credit_amount"]) if t.get("credit_amount") else None,
                "raw_data": {k: str(v) for k, v in t.items() if k not in ("date",)},
            })
        bank_ids = bulk_insert_transactions(report_id, bank_inserts)

        # ── Step 4: Insert ledger transactions ──
        ledger_inserts = []
        for t in ledger_txns_raw:
            amount = t.get("amount")   # None when key absent — falls through to explicit debit/credit fields
            # Tally convention: negative = debit (increase), positive = credit (decrease)
            if amount is not None and isinstance(amount, (int, float, Decimal)):
                amount = float(amount)
                debit = abs(amount) if amount < 0 else None
                credit = amount if amount > 0 else None
            else:
                debit = t.get("debit_amount")
                credit = t.get("credit_amount")

            ledger_inserts.append({
                "source": "ledger",
                "date": _parse_tally_date(t.get("date")),
                "description": t.get("narration") or t.get("description", ""),
                "reference": t.get("voucher_number") or t.get("reference", ""),
                "debit_amount": float(debit) if debit else None,
                "credit_amount": float(credit) if credit else None,
                "voucher_type": t.get("voucher_type", ""),
                "voucher_number": t.get("voucher_number", ""),
                "tally_guid": t.get("guid", ""),
            })
        ledger_ids = bulk_insert_transactions(report_id, ledger_inserts)

        # ── Step 5: Compute closing balances if not provided ──
        if ledger_cb is None and ledger_ob is not None:
            total_debit = sum(t.get("debit_amount") or 0 for t in ledger_inserts)
            total_credit = sum(t.get("credit_amount") or 0 for t in ledger_inserts)
            ledger_cb = ledger_ob + Decimal(str(total_debit)) - Decimal(str(total_credit))

        if bank_cb is None and bank_ob is not None:
            total_credit = sum(t.get("credit_amount") or 0 for t in bank_inserts)
            total_debit = sum(t.get("debit_amount") or 0 for t in bank_inserts)
            bank_cb = bank_ob + Decimal(str(total_credit)) - Decimal(str(total_debit))

        # ── Step 6: Run the shared GROUP matcher ──
        # Adds (over the old 1:1 matcher): reference/UTR matching, a date-proximity
        # window, fuzzy description similarity, one-to-many aggregation, contra/offset
        # detection, and auto/suggest confidence tiers. Build algo dicts directly from
        # inserted data — ledger_ids[i]/bank_ids[i] correspond to ledger_inserts[i]/
        # bank_inserts[i] (bulk_insert_transactions inserts and returns IDs in order).
        from applicationca.agents.document_specific_functions.bank_reconciliation import _to_decimal as _dec
        algo_ledger = [
            {
                "date": t.get("date"),
                "description": t.get("description", ""),
                "reference": t.get("reference", ""),
                "debit_amount": _dec(t.get("debit_amount")),
                "credit_amount": _dec(t.get("credit_amount")),
            }
            for t in ledger_inserts
        ]
        algo_bank = [
            {
                "date": t.get("date"),
                "description": t.get("description", ""),
                "reference": t.get("reference", ""),
                "debit_amount": _dec(t.get("debit_amount")),
                "credit_amount": _dec(t.get("credit_amount")),
            }
            for t in bank_inserts
        ]

        # run_group_matching(ledger_side, bank_side) -> groups whose "book_idxs"
        # index the ledger side and "bank_idxs" the bank side.
        groups = run_group_matching(algo_ledger, algo_bank)

        # ── Step 7: Persist match groups (tier 'auto' -> accepted, 'suggest' -> review) ──
        group_stats = bulk_create_match_groups(
            report_id=report_id,
            groups=groups,
            ledger_ids=ledger_ids,
            bank_ids=bank_ids,
        )

        # ── Step 8: Update report. A group can match many txns at once, so derive the
        #            leftovers from the DB rather than from a 1:1 match count. ──
        total_ledger = len(ledger_inserts)
        total_bank = len(bank_inserts)
        unmatched_ledger = len(get_transactions(report_id, source="ledger", recon_status="unmatched"))
        unmatched_bank = len(get_transactions(report_id, source="bank_statement", recon_status="unmatched"))
        update_report(report_id, user_id, {
            "status": "in_progress",
            "ledger_opening_balance": float(ledger_ob) if ledger_ob else None,
            "ledger_closing_balance": float(ledger_cb) if ledger_cb else None,
            "bank_opening_balance": float(bank_ob) if bank_ob else None,
            "bank_closing_balance": float(bank_cb) if bank_cb else None,
            "total_ledger_transactions": total_ledger,
            "total_bank_transactions": total_bank,
            "matched_count": group_stats["total"],
            "unmatched_ledger_count": unmatched_ledger,
            "unmatched_bank_count": unmatched_bank,
        })

        return {
            "report_id": report_id,
            "status": "in_progress",
            "total_bank_transactions": total_bank,
            "total_ledger_transactions": total_ledger,
            "matched_count": group_stats["total"],
            "auto_matched_groups": group_stats["auto"],
            "suggested_groups": group_stats["suggest"],
            "unmatched_bank_count": unmatched_bank,
            "unmatched_ledger_count": unmatched_ledger,
        }

    except Exception as e:
        logger.error(f"process_report error for report {report_id}: {e}", exc_info=True)
        update_report(report_id, user_id, {"status": "error"})
        return {"error": str(e)}


def _db_row_to_algo(row: Dict) -> Dict:
    """Convert a DB transaction row to the format the matching algorithm expects."""
    from applicationca.agents.document_specific_functions.bank_reconciliation import _to_decimal
    d = row.get("transaction_date")
    return {
        "date": d if isinstance(d, date) else (d.date() if isinstance(d, datetime) else d),
        "description": row.get("description", ""),
        "reference": row.get("reference", ""),
        "debit_amount": _to_decimal(row.get("debit_amount")),
        "credit_amount": _to_decimal(row.get("credit_amount")),
    }


def _parse_tally_date(value) -> Optional[date]:
    """Parse date from Tally (YYYYMMDD or DD-MM-YYYY or date object)."""
    if value is None:
        return None
    if isinstance(value, date):
        return value
    s = str(value).strip()
    # YYYYMMDD
    if len(s) == 8 and s.isdigit():
        try:
            return datetime.strptime(s, "%Y%m%d").date()
        except ValueError:
            pass
    from applicationca.agents.document_specific_functions.bank_reconciliation import _parse_date
    return _parse_date(value)


# ---------------------------------------------------------------------------
# Excel Report Generation
# ---------------------------------------------------------------------------

def generate_excel_report(report_id: int, user_id: str) -> Optional[bytes]:
    """
    Generate an Excel report with:
      Sheet 1: Reconciliation Summary
      Sheet 2: Matched Transactions
      Sheet 3: Unmatched Ledger
      Sheet 4: Unmatched Bank
    Returns bytes of the xlsx file.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.error("openpyxl not installed. Install with: pip install openpyxl")
        return None

    report = get_report(report_id, user_id)
    if not report:
        return None

    summary = get_summary(report_id, user_id)
    wb = openpyxl.Workbook()

    # ── Colors ──
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    subheader_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    total_font = Font(bold=True)
    green_font = Font(color="006400", bold=True)
    red_font = Font(color="8B0000", bold=True)

    def _header_row(ws, values, fill=None, font=None):
        row = ws.max_row + 1 if ws.max_row > 0 else 1
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            if fill:
                cell.fill = fill
            if font:
                cell.font = font
            cell.alignment = Alignment(horizontal="center")

    def _data_row(ws, values):
        row = ws.max_row + 1
        for col, val in enumerate(values, 1):
            ws.cell(row=row, column=col, value=val)

    def _auto_width(ws):
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)

    # ── Sheet 1: Summary ──
    ws1 = wb.active
    ws1.title = "Reconciliation Summary"

    ws1.append([f"Bank Reconciliation — {report['report_name']}"])
    ws1.cell(1, 1).font = Font(bold=True, size=14)
    ws1.append([f"Company: {report['company_name']}   |   Ledger: {report['bank_ledger_name']}"])
    ws1.append([f"Period: {report['period_from']} to {report['period_to']}   |   FY: {report.get('financial_year','')}"])
    ws1.append([])

    ws1.append(["Reconciliation Statement", "", "Amount (₹)"])
    ws1.cell(ws1.max_row, 1).font = Font(bold=True, size=12)

    def _summary_row(label, amount, sign=""):
        ws1.append([label, sign, f"{amount:,.2f}"])

    lcb = summary["ledger_closing_balance"]
    bcb = summary["bank_closing_balance"]
    s = summary["summary"]

    _summary_row("Ledger Closing Balance", lcb)
    _summary_row("Add: Ledger Credits not in Bank Statement", s["ledger_credits_not_in_bank"], "(+)")
    _summary_row("Less: Ledger Debits not in Bank Statement", s["ledger_debits_not_in_bank"], "(-)")
    _summary_row("Add: Bank Credits not in Ledger", s["bank_credits_not_in_ledger"], "(+)")
    _summary_row("Less: Bank Debits not in Ledger", s["bank_debits_not_in_ledger"], "(-)")
    ws1.append(["─" * 40])
    _summary_row("Reconciled Amount", s["reconciled_amount"])
    _summary_row("Bank Statement Closing Balance", bcb)
    ws1.append(["─" * 40])
    diff_row = ws1.max_row + 1
    ws1.append(["Reconciliation Difference", "", f"{s['reconciliation_difference']:,.2f}"])
    diff_cell = ws1.cell(diff_row, 3)
    diff_val = s["reconciliation_difference"]
    diff_cell.font = green_font if abs(diff_val) < 0.01 else red_font

    _auto_width(ws1)

    # ── Sheet 2: Matched Transactions ──
    ws2 = wb.create_sheet("Matched Transactions")
    _header_row(ws2, [
        "Match Type", "Status",
        "Bank Date", "Bank Description", "Bank Ref", "Bank Dr", "Bank Cr",
        "Ledger Date", "Ledger Description", "Ledger Ref", "Ledger Dr", "Ledger Cr",
        "Voucher Type", "Voucher No",
    ], fill=header_fill, font=header_font)

    for m in get_matches(report_id, status="accepted"):
        banks = m.get("bank_transactions") or []
        ledgers = m.get("ledger_transactions") or []
        # One row per member pair; type/status printed only on the first line of a group.
        for idx in range(max(len(banks), len(ledgers), 1)):
            bt = banks[idx] if idx < len(banks) else {}
            lt = ledgers[idx] if idx < len(ledgers) else {}
            _data_row(ws2, [
                m["match_type"] if idx == 0 else "", m["status"] if idx == 0 else "",
                bt.get("transaction_date"), bt.get("description"), bt.get("reference"),
                bt.get("debit_amount"), bt.get("credit_amount"),
                lt.get("transaction_date"), lt.get("description"), lt.get("reference"),
                lt.get("debit_amount"), lt.get("credit_amount"),
                lt.get("voucher_type"), lt.get("voucher_number"),
            ])
    _auto_width(ws2)

    # ── Sheet 3: Unmatched Ledger ──
    ws3 = wb.create_sheet("Unmatched Ledger Entries")
    _header_row(ws3, ["Date", "Description", "Reference", "Debit", "Credit", "Voucher Type", "Voucher No"],
                fill=header_fill, font=header_font)
    for t in get_transactions(report_id, source="ledger", recon_status="unmatched"):
        _data_row(ws3, [
            t["transaction_date"], t["description"], t["reference"],
            t["debit_amount"], t["credit_amount"],
            t["voucher_type"], t["voucher_number"],
        ])
    _auto_width(ws3)

    # ── Sheet 4: Unmatched Bank ──
    ws4 = wb.create_sheet("Unmatched Bank Entries")
    _header_row(ws4, ["Date", "Description", "Reference", "Debit", "Credit"],
                fill=header_fill, font=header_font)
    for t in get_transactions(report_id, source="bank_statement", recon_status="unmatched"):
        _data_row(ws4, [
            t["transaction_date"], t["description"], t["reference"],
            t["debit_amount"], t["credit_amount"],
        ])
    _auto_width(ws4)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
