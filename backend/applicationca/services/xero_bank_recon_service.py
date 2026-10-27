"""
Xero Bank Reconciliation Service.

Clone of `qb_bank_recon_service` (itself a clone of the Tally `bank_recon_service`
flow) adapted for Xero so Xero walks the exact same stateful path:

    create report (persist statement)
      -> process (parse statement + fetch Xero txns + derive closing balance
                  + run the shared auto-matcher + persist normalized rows)
      -> review matches (accept / reject / manual)
      -> summary -> save -> Excel

The only ERP-specific swaps vs Tally/QBO are the "book" side fetch (Xero
BankTransactions instead of the desktop bridge / QBO TransactionList) and the
book closing balance (derived from Xero's account balance). The statement
parser, matcher and summary math are the shared ones, so Xero, QB and Tally
reconcile to the same numbers.
"""

import io
import json
import logging
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Dict, List, Optional

import boto3

from config.settings import settings
from applicationca.database.database_retrieval import (
    execute_query,
    get_connection,
    return_connection,
)
# Reuse the Tally service's statement-file storage helpers verbatim — there's
# nothing ERP-specific about staging an uploaded PDF to S3 / memory.
from applicationca.services.bank_recon_service import (
    _upload_to_s3,
    store_pending_file,
    pop_pending_file,
)
from applicationca.services.xero_bank import (
    get_account_closing_balance,
    list_bank_txns,
)
from applicationca.agents.document_specific_functions.bank_reconciliation import (
    compute_reconciliation_summary,
    get_financial_year,
    parse_bank_statement_with_llm,
    run_automatic_matching,
    _to_decimal as _dec,
)

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Serial number generation
# ---------------------------------------------------------------------------

def get_next_report_serial(user_id: str, today: date) -> str:
    """Return next serial string like 'XeroBankRecon_20260608_001'."""
    date_str = today.strftime("%Y%m%d")
    prefix = f"XeroBankRecon_{date_str}_"
    rows = execute_query(
        """
        SELECT report_name FROM xero_bank_recon_reports
        WHERE user_id = %s AND report_name LIKE %s
        ORDER BY report_name DESC LIMIT 1
        """,
        (user_id, f"{prefix}%"),
    )
    if rows:
        last_name = rows[0]["report_name"]
        try:
            last_num = int(last_name.split("_")[-1])
        except (ValueError, IndexError):
            last_num = 0
        return f"{prefix}{str(last_num + 1).zfill(3)}"
    return f"{prefix}001"


# ---------------------------------------------------------------------------
# Report CRUD
# ---------------------------------------------------------------------------

def create_report(
    user_id: str,
    tenant_id: str,
    account_id: str,
    account_name: Optional[str],
    report_name: str,
    period_from: date,
    period_to: date,
    financial_year: str,
    file_url: str,
    file_name: str,
    xero_closing_balance_override: Optional[float] = None,
) -> Optional[Dict]:
    conn = get_connection()
    if not conn:
        return None
    try:
        cursor = conn.cursor()
        cursor.execute(
            """
            INSERT INTO xero_bank_recon_reports
              (user_id, tenant_id, account_id, account_name, report_name,
               period_from, period_to, financial_year,
               status, xero_closing_balance_override,
               bank_statement_file_url, bank_statement_file_name)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,'draft',%s,%s,%s)
            RETURNING report_id, report_name, status, created_at
            """,
            (user_id, tenant_id, account_id, account_name, report_name,
             period_from, period_to, financial_year,
             xero_closing_balance_override, file_url, file_name),
        )
        row = cursor.fetchone()
        conn.commit()
        return {
            "report_id": row[0],
            "report_name": row[1],
            "status": row[2],
            "created_at": row[3].isoformat() if row[3] else None,
        }
    except Exception as e:
        conn.rollback()
        logger.error(f"xero create_report error: {e}")
        return None
    finally:
        return_connection(conn)


def get_reports(user_id: str, tenant_id: str, account_id: str) -> List[Dict]:
    rows = execute_query(
        """
        SELECT report_id, report_name, tenant_id, account_id, account_name,
               period_from, period_to, financial_year, status,
               book_opening_balance, bank_opening_balance,
               book_closing_balance, bank_closing_balance,
               reconciliation_difference, matched_count,
               unmatched_book_count, unmatched_bank_count,
               created_at, updated_at
        FROM xero_bank_recon_reports
        WHERE user_id = %s AND tenant_id = %s AND account_id = %s
        ORDER BY created_at DESC
        """,
        (user_id, tenant_id, account_id),
    )
    return [dict(r) for r in rows] if rows else []


def get_report(report_id: int, user_id: str) -> Optional[Dict]:
    rows = execute_query(
        """
        SELECT * FROM xero_bank_recon_reports
        WHERE report_id = %s AND user_id = %s
        """,
        (report_id, user_id),
    )
    return dict(rows[0]) if rows else None


def update_report(report_id: int, user_id: str, fields: Dict) -> bool:
    allowed = {
        "report_name", "status", "account_name",
        "book_opening_balance", "book_closing_balance",
        "bank_opening_balance", "bank_closing_balance",
        "xero_closing_balance_override",
        "reconciliation_difference",
        "book_credits_not_in_bank", "book_debits_not_in_bank",
        "bank_credits_not_in_book", "bank_debits_not_in_book",
        "total_book_transactions", "total_bank_transactions",
        "matched_count", "unmatched_book_count", "unmatched_bank_count",
    }
    update_fields = {k: v for k, v in fields.items() if k in allowed}
    if not update_fields:
        return False
    set_clause = ", ".join(f"{k} = %s" for k in update_fields)
    values = list(update_fields.values()) + [datetime.now(), report_id, user_id]
    result = execute_query(
        f"""
        UPDATE xero_bank_recon_reports
        SET {set_clause}, updated_at = %s
        WHERE report_id = %s AND user_id = %s
        """,
        values,
    )
    return result is not None


def delete_report(report_id: int, user_id: str) -> bool:
    result = execute_query(
        "DELETE FROM xero_bank_recon_reports WHERE report_id = %s AND user_id = %s",
        (report_id, user_id),
    )
    return result is not None


# ---------------------------------------------------------------------------
# Transaction CRUD
# ---------------------------------------------------------------------------

def bulk_insert_transactions(report_id: int, transactions: List[Dict]) -> List[int]:
    """Insert many transactions in a single batch and return their txn_ids in order."""
    if not transactions:
        return []
    conn = get_connection()
    if not conn:
        return []
    try:
        from psycopg2.extras import execute_values
        cursor = conn.cursor()
        rows = [
            (
                report_id,
                t["source"],
                t.get("date"),
                t.get("description", ""),
                t.get("reference", ""),
                t.get("debit_amount"),
                t.get("credit_amount"),
                t.get("xero_txn_id"),
                t.get("xero_txn_type"),
                json.dumps(t.get("raw_data")) if t.get("raw_data") else None,
                t.get("is_opening_balance", False),
            )
            for t in transactions
        ]
        fetched = execute_values(
            cursor,
            """
            INSERT INTO xero_bank_recon_transactions
              (report_id, source, transaction_date, description, reference,
               debit_amount, credit_amount, recon_status,
               xero_txn_id, xero_txn_type, raw_data, is_opening_balance)
            VALUES %s
            RETURNING txn_id
            """,
            rows,
            template="(%s,%s,%s,%s,%s,%s,%s,'unmatched',%s,%s,%s,%s)",
            fetch=True,
        )
        ids = [row[0] for row in fetched]
        conn.commit()
        return ids
    except Exception as e:
        conn.rollback()
        logger.error(f"xero bulk_insert_transactions error: {e}")
        return []
    finally:
        return_connection(conn)


def get_transactions(
    report_id: int,
    source: Optional[str] = None,
    recon_status: Optional[str] = None,
) -> List[Dict]:
    base_sql = """
        SELECT txn_id, report_id, source, transaction_date, description, reference,
               debit_amount, credit_amount, recon_status, match_id,
               xero_txn_id, xero_txn_type, is_opening_balance, created_at
        FROM xero_bank_recon_transactions
        WHERE report_id = %s AND is_opening_balance = FALSE
    """
    params = [report_id]
    if source:
        base_sql += " AND source = %s"
        params.append(source)
    if recon_status:
        base_sql += " AND recon_status = %s"
        params.append(recon_status)
    base_sql += " ORDER BY transaction_date, txn_id"
    rows = execute_query(base_sql, params)
    return [dict(r) for r in rows] if rows else []


def mark_transactions_pending(report_id: int, txn_ids: List[int], pending: bool = True) -> int:
    """Move unmatched transactions to/from the 'pending' triage bucket.

    No Xero writes — this is pure recon triage ("set this row aside for later").
    A pending row drops out of the unmatched lists and the reconciliation
    difference. Only `unmatched` <-> `pending` transitions are allowed, so
    matched rows are never disturbed. Returns the number of rows changed.
    """
    if not txn_ids:
        return 0
    target = "pending" if pending else "unmatched"
    source = "unmatched" if pending else "pending"
    conn = get_connection()
    if not conn:
        return 0
    try:
        cursor = conn.cursor()
        placeholders = ",".join(["%s"] * len(txn_ids))
        cursor.execute(
            f"""
            UPDATE xero_bank_recon_transactions
            SET recon_status = %s
            WHERE report_id = %s AND recon_status = %s AND txn_id IN ({placeholders})
            """,
            [target, report_id, source, *txn_ids],
        )
        n = cursor.rowcount
        conn.commit()
        return n
    except Exception as e:
        conn.rollback()
        logger.error(f"xero mark_transactions_pending error: {e}")
        return 0
    finally:
        return_connection(conn)


# ---------------------------------------------------------------------------
# Match CRUD
# ---------------------------------------------------------------------------

def bulk_create_matches(
    report_id: int,
    match_results: List[Dict],
    book_ids: List[int],
    bank_ids: List[int],
    match_type: str = "automatic",
) -> int:
    """Insert all matches and update transaction statuses in one DB transaction.

    `match_results` come from run_automatic_matching: each has ledger_idx /
    bank_idx pointing into the book / bank inserted-row order, so book_ids[idx]
    and bank_ids[idx] recover the persisted txn_ids.
    """
    if not match_results:
        return 0
    conn = get_connection()
    if not conn:
        return 0
    try:
        from psycopg2.extras import execute_values
        cursor = conn.cursor()

        rows = [
            (
                report_id,
                match_type,
                book_ids[m["ledger_idx"]],
                bank_ids[m["bank_idx"]],
                m["confidence"],
                m["reason"],
                "",
            )
            for m in match_results
        ]
        inserted = execute_values(
            cursor,
            """
            INSERT INTO xero_bank_recon_matches
              (report_id, match_type, status, book_txn_id, bank_txn_id,
               confidence_score, match_reason, notes)
            VALUES %s
            RETURNING match_id, book_txn_id, bank_txn_id
            """,
            rows,
            template="(%s,%s,'pending_review',%s,%s,%s,%s,%s)",
            fetch=True,
        )

        new_status = "auto_matched" if match_type == "automatic" else "manual_matched"
        status_updates = [
            (new_status, mid, txn_id)
            for mid, btid, bktid in inserted
            for txn_id in (btid, bktid)
        ]
        cursor.executemany(
            "UPDATE xero_bank_recon_transactions SET recon_status = %s, match_id = %s WHERE txn_id = %s",
            status_updates,
        )

        conn.commit()
        return len(inserted)
    except Exception as e:
        conn.rollback()
        logger.error(f"xero bulk_create_matches error: {e}")
        return 0
    finally:
        return_connection(conn)


def create_match(
    report_id: int,
    book_txn_id: int,
    bank_txn_id: int,
    match_type: str,
    confidence: float,
    reason: str,
    notes: str = "",
) -> Optional[int]:
    conn = get_connection()
    if not conn:
        return None
    try:
        cursor = conn.cursor()
        cursor.execute(
            """
            INSERT INTO xero_bank_recon_matches
              (report_id, match_type, status, book_txn_id, bank_txn_id,
               confidence_score, match_reason, notes)
            VALUES (%s,%s,'pending_review',%s,%s,%s,%s,%s)
            RETURNING match_id
            """,
            (report_id, match_type, book_txn_id, bank_txn_id, confidence, reason, notes),
        )
        match_id = cursor.fetchone()[0]
        conn.commit()
        return match_id
    except Exception as e:
        conn.rollback()
        logger.error(f"xero create_match error: {e}")
        return None
    finally:
        return_connection(conn)


def get_matches(report_id: int, status: Optional[str] = None) -> List[Dict]:
    sql = """
        SELECT m.match_id, m.report_id, m.match_type, m.status,
               m.confidence_score, m.match_reason, m.notes,
               m.created_at, m.reviewed_at,
               -- bank transaction
               bt.txn_id AS bank_txn_id, bt.transaction_date AS bank_date,
               bt.description AS bank_desc, bt.reference AS bank_ref,
               bt.debit_amount AS bank_debit, bt.credit_amount AS bank_credit,
               -- book (Xero) transaction
               kt.txn_id AS book_txn_id, kt.transaction_date AS book_date,
               kt.description AS book_desc, kt.reference AS book_ref,
               kt.debit_amount AS book_debit, kt.credit_amount AS book_credit,
               kt.xero_txn_id, kt.xero_txn_type
        FROM xero_bank_recon_matches m
        JOIN xero_bank_recon_transactions bt ON bt.txn_id = m.bank_txn_id
        JOIN xero_bank_recon_transactions kt ON kt.txn_id = m.book_txn_id
        WHERE m.report_id = %s
    """
    params = [report_id]
    if status:
        sql += " AND m.status = %s"
        params.append(status)
    sql += " ORDER BY m.created_at"
    rows = execute_query(sql, params)
    if not rows:
        return []

    result = []
    for r in rows:
        result.append({
            "match_id": r["match_id"],
            "match_type": r["match_type"],
            "status": r["status"],
            "confidence_score": float(r["confidence_score"]) if r["confidence_score"] else None,
            "match_reason": r["match_reason"],
            "notes": r["notes"],
            "created_at": r["created_at"].isoformat() if r["created_at"] else None,
            "reviewed_at": r["reviewed_at"].isoformat() if r["reviewed_at"] else None,
            "bank_transaction": {
                "txn_id": r["bank_txn_id"],
                "transaction_date": r["bank_date"].isoformat() if r["bank_date"] else None,
                "description": r["bank_desc"],
                "reference": r["bank_ref"],
                "debit_amount": float(r["bank_debit"]) if r["bank_debit"] else None,
                "credit_amount": float(r["bank_credit"]) if r["bank_credit"] else None,
            },
            "book_transaction": {
                "txn_id": r["book_txn_id"],
                "transaction_date": r["book_date"].isoformat() if r["book_date"] else None,
                "description": r["book_desc"],
                "reference": r["book_ref"],
                "debit_amount": float(r["book_debit"]) if r["book_debit"] else None,
                "credit_amount": float(r["book_credit"]) if r["book_credit"] else None,
                "xero_txn_id": r["xero_txn_id"],
                "xero_txn_type": r["xero_txn_type"],
            },
        })
    return result


def update_match_status(match_id: int, status: str) -> Optional[Dict]:
    """Accept or reject a match. On reject, return both txns to unmatched."""
    conn = get_connection()
    if not conn:
        logger.error("xero update_match_status: could not get DB connection")
        return None
    try:
        cursor = conn.cursor()
        cursor.execute(
            """
            UPDATE xero_bank_recon_matches
            SET status = %s
            WHERE match_id = %s
            RETURNING match_id, book_txn_id, bank_txn_id, match_type
            """,
            (status, match_id),
        )
        row = cursor.fetchone()
        if not row:
            conn.rollback()
            logger.warning(f"xero update_match_status: match_id={match_id} not found")
            return None

        mid, book_txn_id, bank_txn_id, match_type = row

        if status == "accepted":
            new_status = "auto_matched" if match_type == "automatic" else "manual_matched"
            cursor.execute(
                "UPDATE xero_bank_recon_transactions SET recon_status = %s, match_id = %s WHERE txn_id = %s",
                (new_status, mid, book_txn_id),
            )
            cursor.execute(
                "UPDATE xero_bank_recon_transactions SET recon_status = %s, match_id = %s WHERE txn_id = %s",
                (new_status, mid, bank_txn_id),
            )
        elif status == "rejected":
            cursor.execute(
                "UPDATE xero_bank_recon_transactions SET recon_status = 'unmatched', match_id = NULL WHERE txn_id IN (%s, %s)",
                (book_txn_id, bank_txn_id),
            )

        conn.commit()

        # Best-effort reviewed_at stamp
        try:
            cursor2 = conn.cursor()
            cursor2.execute(
                "UPDATE xero_bank_recon_matches SET reviewed_at = NOW() WHERE match_id = %s",
                (mid,),
            )
            conn.commit()
        except Exception:
            conn.rollback()

        return {"match_id": mid, "status": status}
    except Exception as e:
        conn.rollback()
        logger.error(f"xero update_match_status error (match_id={match_id}, status={status}): {e}")
        return None
    finally:
        return_connection(conn)


def bulk_update_match_status(match_ids: List[int], status: str) -> int:
    updated = 0
    for mid in match_ids:
        if update_match_status(mid, status):
            updated += 1
    return updated


def create_manual_match(
    report_id: int,
    book_txn_id: int,
    bank_txn_id: int,
    notes: str = "",
) -> Optional[Dict]:
    match_id = create_match(
        report_id=report_id,
        book_txn_id=book_txn_id,
        bank_txn_id=bank_txn_id,
        match_type="manual",
        confidence=1.0,
        reason="Manual match by user",
        notes=notes,
    )
    if not match_id:
        return None
    return update_match_status(match_id, "accepted")


def delete_match(match_id: int) -> bool:
    """Remove a match and return both transactions to unmatched state."""
    conn = get_connection()
    if not conn:
        return False
    try:
        cursor = conn.cursor()
        cursor.execute(
            "SELECT book_txn_id, bank_txn_id FROM xero_bank_recon_matches WHERE match_id = %s",
            (match_id,),
        )
        row = cursor.fetchone()
        if not row:
            return False
        book_txn_id, bank_txn_id = row
        cursor.execute("DELETE FROM xero_bank_recon_matches WHERE match_id = %s", (match_id,))
        cursor.execute(
            "UPDATE xero_bank_recon_transactions SET recon_status='unmatched', match_id=NULL WHERE txn_id IN (%s,%s)",
            (book_txn_id, bank_txn_id),
        )
        conn.commit()
        return True
    except Exception as e:
        conn.rollback()
        logger.error(f"xero delete_match error: {e}")
        return False
    finally:
        return_connection(conn)


# ---------------------------------------------------------------------------
# Reconciliation Summary
# ---------------------------------------------------------------------------

def _effective_book_cb(report: Dict) -> Decimal:
    """User override wins over the auto-derived book closing balance."""
    override = report.get("xero_closing_balance_override")
    if override is not None:
        return Decimal(str(override))
    return Decimal(str(report.get("book_closing_balance") or 0))


def get_summary(report_id: int, user_id: str) -> Optional[Dict]:
    report = get_report(report_id, user_id)
    if not report:
        return None

    unmatched_book = get_transactions(report_id, source="book", recon_status="unmatched")
    unmatched_bank = get_transactions(report_id, source="bank_statement", recon_status="unmatched")

    book_cb = _effective_book_cb(report)
    bank_cb = Decimal(str(report.get("bank_closing_balance") or 0))

    summary = compute_reconciliation_summary(book_cb, bank_cb, unmatched_book, unmatched_bank)

    pending_count = execute_query(
        "SELECT COUNT(*) as cnt FROM xero_bank_recon_matches WHERE report_id = %s AND status = 'pending_review'",
        (report_id,),
    )
    accepted_count = execute_query(
        "SELECT COUNT(*) as cnt FROM xero_bank_recon_matches WHERE report_id = %s AND status = 'accepted'",
        (report_id,),
    )
    pending_txn_count = execute_query(
        "SELECT COUNT(*) as cnt FROM xero_bank_recon_transactions WHERE report_id = %s AND recon_status = 'pending'",
        (report_id,),
    )

    return {
        "report_id": report_id,
        "report_name": report.get("report_name"),
        "status": report.get("status"),
        "tenant_id": report.get("tenant_id"),
        "account_id": report.get("account_id"),
        "account_name": report.get("account_name"),
        "period_from": report["period_from"].isoformat() if report.get("period_from") else None,
        "period_to": report["period_to"].isoformat() if report.get("period_to") else None,
        "financial_year": report.get("financial_year"),
        "book_opening_balance": float(report.get("book_opening_balance") or 0),
        "book_closing_balance": float(book_cb),
        "bank_opening_balance": float(report.get("bank_opening_balance") or 0),
        "bank_closing_balance": float(bank_cb),
        "summary": summary,
        "counts": {
            "total_matched": report.get("matched_count", 0),
            "unmatched_book": len(unmatched_book),
            "unmatched_bank": len(unmatched_bank),
            "pending_review": pending_count[0]["cnt"] if pending_count else 0,
            "accepted": accepted_count[0]["cnt"] if accepted_count else 0,
            "pending_transactions": pending_txn_count[0]["cnt"] if pending_txn_count else 0,
        },
    }


# ---------------------------------------------------------------------------
# Processing Pipeline
# ---------------------------------------------------------------------------

def _xero_txn_description(t: Dict) -> str:
    return " ".join(s for s in (t.get("payee"), t.get("memo"), t.get("split_account")) if s)


async def process_report(report_id: int, user_id: str, tenant_id: Optional[str]) -> Dict:
    """
    Full processing pipeline (Xero analogue of qb_bank_recon_service.process_report):
      1. Load + parse the saved bank statement
      2. Fetch Xero transactions for the account in the period (BankTransactions)
      3. Derive the book closing balance as of period end
      4. Insert both sides into xero_bank_recon_transactions
      5. Run the shared automatic matcher
      6. Persist matches (pending_review) + update report counts/balances
    """
    report = get_report(report_id, user_id)
    if not report:
        return {"error": "Report not found"}

    tenant_id = tenant_id or report.get("tenant_id")

    # Idempotency: don't reprocess an in-flight/finished report (frontend retries).
    # awaiting_source falls through — it means "ready to retry now we're online".
    current_status = report.get("status")
    if current_status in ("processing", "completed"):
        existing_matches = get_matches(report_id)
        return {
            "status": current_status,
            "message": f"Report already {current_status}",
            "matches_count": len(existing_matches),
            "idempotent": True,
        }

    update_report(report_id, user_id, {"status": "processing"})

    try:
        # ── Step 1: load the bank statement file ──
        file_url = report.get("bank_statement_file_url", "")
        file_name = report.get("bank_statement_file_name", "statement.pdf")
        file_content: Optional[bytes] = None

        if file_url and file_url.startswith("pending://"):
            pending = pop_pending_file(file_url)
            if pending:
                file_content, file_name = pending
            else:
                update_report(report_id, user_id, {"status": "error"})
                return {"error": "Uploaded file not found in memory. Please re-upload the statement."}
        elif file_url:
            try:
                s3 = boto3.client(
                    "s3",
                    aws_access_key_id=settings.AWS_ACCESS_KEY_ID,
                    aws_secret_access_key=settings.AWS_SECRET_ACCESS_KEY,
                    region_name=settings.AWS_DEFAULT_REGION,
                )
                bucket = settings.AWS_BUCKET_NAME if hasattr(settings, "AWS_BUCKET_NAME") else "aureus-documents"
                s3_key = file_url.split(f"{bucket}.s3.{settings.AWS_DEFAULT_REGION}.amazonaws.com/")[-1]
                obj = s3.get_object(Bucket=bucket, Key=s3_key)
                file_content = obj["Body"].read()
            except Exception as e:
                logger.error(f"xero S3 download failed: {e}")
                update_report(report_id, user_id, {"status": "error"})
                return {"error": "Could not retrieve bank statement file from storage."}
        else:
            update_report(report_id, user_id, {"status": "error"})
            return {"error": "No bank statement file associated with this report."}

        period_from: date = report["period_from"]
        period_to: date = report["period_to"]
        account_id = report["account_id"]

        # ── Step 2: parse statement + fetch Xero book txns ──
        # Use the SAME LLM parser Tally recon uses (parse_bank_statement_with_llm),
        # not the rule-based parser. The rule-based path can't tell debit from
        # credit on text-only PDFs (it dumps everything into debit_amount), which
        # silently breaks the cross-amount matcher. The LLM parser also filters
        # rows to the recon period. Same engine as Tally => same matches.
        parsed = await parse_bank_statement_with_llm(
            file_content, file_name, period_from=period_from, period_to=period_to,
        )
        bank_txns_raw = [
            t for t in (parsed.get("transactions") or [])
            if not t.get("is_opening_balance") and not t.get("is_closing_balance")
        ]
        bank_ob = parsed.get("opening_balance")
        bank_cb = parsed.get("closing_balance")

        xero_txns = await list_bank_txns(
            user_id,
            account_id=account_id,
            start_date=period_from.isoformat(),
            end_date=period_to.isoformat(),
            tenant_id=tenant_id,
        )
        logger.info(
            f"xero recon report {report_id}: parsed {len(bank_txns_raw)} statement rows, "
            f"fetched {len(xero_txns)} Xero txns"
        )

        # ── Step 3: derive book closing balance (override wins) ──
        override = report.get("xero_closing_balance_override")
        if override is not None:
            book_cb = Decimal(str(override))
        else:
            book_cb = await get_account_closing_balance(
                user_id, account_id=account_id, as_of=period_to.isoformat(), tenant_id=tenant_id
            )

        # ── Step 4a: insert Xero book transactions (inflow→debit, outflow→credit) ──
        book_inserts = []
        for t in xero_txns:
            amount = t.get("amount") or 0.0
            debit = abs(amount) if amount > 0 else None
            credit = abs(amount) if amount < 0 else None
            tdate = _parse_iso_date(t.get("date"))
            if not tdate:
                continue
            book_inserts.append({
                "source": "book",
                "date": tdate,
                "description": _xero_txn_description(t),
                "reference": t.get("doc_number") or "",
                "debit_amount": float(debit) if debit else None,
                "credit_amount": float(credit) if credit else None,
                "xero_txn_id": t.get("id"),
                "xero_txn_type": t.get("type"),
                "raw_data": {k: (str(v) if v is not None else None) for k, v in t.items()},
            })
        book_ids = bulk_insert_transactions(report_id, book_inserts)

        # ── Step 4b: insert bank statement transactions ──
        bank_inserts = []
        for t in bank_txns_raw:
            tdate = _coerce_date(t.get("date"))
            if not tdate:
                continue
            bank_inserts.append({
                "source": "bank_statement",
                "date": tdate,
                "description": t.get("description", ""),
                "reference": t.get("reference", "") or "",
                "debit_amount": float(t["debit_amount"]) if t.get("debit_amount") else None,
                "credit_amount": float(t["credit_amount"]) if t.get("credit_amount") else None,
                "raw_data": {k: str(v) for k, v in t.items() if k not in ("date",)},
            })
        bank_ids = bulk_insert_transactions(report_id, bank_inserts)

        # ── Step 5: run the shared automatic matcher ──
        algo_book = [
            {
                "date": t["date"],
                "description": t.get("description", ""),
                "reference": t.get("reference", ""),
                "debit_amount": _dec(t.get("debit_amount")),
                "credit_amount": _dec(t.get("credit_amount")),
            }
            for t in book_inserts
        ]
        algo_bank = [
            {
                "date": t["date"],
                "description": t.get("description", ""),
                "reference": t.get("reference", ""),
                "debit_amount": _dec(t.get("debit_amount")),
                "credit_amount": _dec(t.get("credit_amount")),
            }
            for t in bank_inserts
        ]
        match_results = run_automatic_matching(algo_book, algo_bank)

        # ── Step 6: persist matches + update report ──
        match_count = bulk_create_matches(
            report_id=report_id,
            match_results=match_results,
            book_ids=book_ids,
            bank_ids=bank_ids,
            match_type="automatic",
        )

        total_book = len(book_inserts)
        total_bank = len(bank_inserts)
        update_report(report_id, user_id, {
            "status": "in_progress",
            "book_closing_balance": float(book_cb) if book_cb is not None else None,
            "bank_opening_balance": float(bank_ob) if bank_ob is not None else None,
            "bank_closing_balance": float(bank_cb) if bank_cb is not None else None,
            "total_book_transactions": total_book,
            "total_bank_transactions": total_bank,
            "matched_count": match_count,
            "unmatched_book_count": total_book - match_count,
            "unmatched_bank_count": total_bank - match_count,
        })

        return {
            "report_id": report_id,
            "status": "in_progress",
            "total_bank_transactions": total_bank,
            "total_book_transactions": total_book,
            "matched_count": match_count,
            "unmatched_bank_count": total_bank - match_count,
            "unmatched_book_count": total_book - match_count,
        }

    except Exception as e:
        logger.error(f"xero process_report error for report {report_id}: {e}", exc_info=True)
        update_report(report_id, user_id, {"status": "error"})
        return {"error": str(e)}


def _parse_iso_date(value) -> Optional[date]:
    """Parse a Xero 'YYYY-MM-DD' date string (or date object) to a date."""
    if value is None:
        return None
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
    except (ValueError, TypeError):
        return None


def _coerce_date(value) -> Optional[date]:
    """Coerce a parsed statement-row date (date/datetime/str) to a date."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    from applicationca.agents.document_specific_functions.bank_reconciliation import _parse_date
    return _parse_date(value)


# ---------------------------------------------------------------------------
# Excel Report Generation
# ---------------------------------------------------------------------------

def generate_excel_report(report_id: int, user_id: str) -> Optional[bytes]:
    """Excel export mirroring the Tally / QBO recon report:
      Sheet 1: Reconciliation Summary
      Sheet 2: Matched Transactions
      Sheet 3: Unmatched Book (Xero)
      Sheet 4: Unmatched Bank
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.error("openpyxl not installed. Install with: pip install openpyxl")
        return None

    report = get_report(report_id, user_id)
    if not report:
        return None

    summary = get_summary(report_id, user_id)
    wb = openpyxl.Workbook()

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    green_font = Font(color="006400", bold=True)
    red_font = Font(color="8B0000", bold=True)

    def _header_row(ws, values):
        row = ws.max_row + 1 if ws.max_row > 0 else 1
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

    def _data_row(ws, values):
        row = ws.max_row + 1
        for col, val in enumerate(values, 1):
            ws.cell(row=row, column=col, value=val)

    def _auto_width(ws):
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)

    # ── Sheet 1: Summary ──
    ws1 = wb.active
    ws1.title = "Reconciliation Summary"
    ws1.append([f"Bank Reconciliation — {report['report_name']}"])
    ws1.cell(1, 1).font = Font(bold=True, size=14)
    ws1.append([f"Account: {report.get('account_name') or report.get('account_id')}"])
    ws1.append([f"Period: {report['period_from']} to {report['period_to']}   |   FY: {report.get('financial_year','')}"])
    ws1.append([])

    s = summary["summary"]
    bcb = summary["bank_closing_balance"]
    kcb = summary["book_closing_balance"]

    def _summary_row(label, amount, sign=""):
        ws1.append([label, sign, f"{amount:,.2f}"])

    ws1.append(["Reconciliation Statement", "", "Amount"])
    ws1.cell(ws1.max_row, 1).font = Font(bold=True, size=12)
    _summary_row("Xero Book Closing Balance", kcb)
    _summary_row("Add: Book Credits not in Bank Statement", s["ledger_credits_not_in_bank"], "(+)")
    _summary_row("Less: Book Debits not in Bank Statement", s["ledger_debits_not_in_bank"], "(-)")
    _summary_row("Add: Bank Credits not in Book", s["bank_credits_not_in_ledger"], "(+)")
    _summary_row("Less: Bank Debits not in Book", s["bank_debits_not_in_ledger"], "(-)")
    _summary_row("Reconciled Amount", s["reconciled_amount"])
    _summary_row("Bank Statement Closing Balance", bcb)
    diff_row = ws1.max_row + 1
    ws1.append(["Reconciliation Difference", "", f"{s['reconciliation_difference']:,.2f}"])
    diff_cell = ws1.cell(diff_row, 3)
    diff_cell.font = green_font if abs(s["reconciliation_difference"]) < 0.01 else red_font
    _auto_width(ws1)

    # ── Sheet 2: Matched ──
    ws2 = wb.create_sheet("Matched Transactions")
    _header_row(ws2, [
        "Match Type", "Status",
        "Bank Date", "Bank Description", "Bank Ref", "Bank Dr", "Bank Cr",
        "Book Date", "Book Description", "Book Ref", "Book Dr", "Book Cr",
        "Xero Txn Type", "Xero Txn ID",
    ])
    for m in get_matches(report_id, status="accepted"):
        bt = m["bank_transaction"]
        kt = m["book_transaction"]
        _data_row(ws2, [
            m["match_type"], m["status"],
            bt["transaction_date"], bt["description"], bt["reference"],
            bt["debit_amount"], bt["credit_amount"],
            kt["transaction_date"], kt["description"], kt["reference"],
            kt["debit_amount"], kt["credit_amount"],
            kt["xero_txn_type"], kt["xero_txn_id"],
        ])
    _auto_width(ws2)

    # ── Sheet 3: Unmatched Book ──
    ws3 = wb.create_sheet("Unmatched Book Entries")
    _header_row(ws3, ["Date", "Description", "Reference", "Debit", "Credit", "Xero Txn Type", "Xero Txn ID"])
    for t in get_transactions(report_id, source="book", recon_status="unmatched"):
        _data_row(ws3, [
            t["transaction_date"], t["description"], t["reference"],
            t["debit_amount"], t["credit_amount"], t["xero_txn_type"], t["xero_txn_id"],
        ])
    _auto_width(ws3)

    # ── Sheet 4: Unmatched Bank ──
    ws4 = wb.create_sheet("Unmatched Bank Entries")
    _header_row(ws4, ["Date", "Description", "Reference", "Debit", "Credit"])
    for t in get_transactions(report_id, source="bank_statement", recon_status="unmatched"):
        _data_row(ws4, [
            t["transaction_date"], t["description"], t["reference"],
            t["debit_amount"], t["credit_amount"],
        ])
    _auto_width(ws4)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
