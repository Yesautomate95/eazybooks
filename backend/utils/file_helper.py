import io
import os

import PyPDF2
import docx
import openpyxl
from google.cloud import vision
from google.oauth2 import service_account

credentials = service_account.Credentials.from_service_account_file(
    "data/deft-set-448318-p4-dc4d486cfd82.json"
)

# Initialize the Vision API client with the credentials
vision_client = vision.ImageAnnotatorClient(credentials=credentials)

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF using PyPDF2."""
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        text = ""
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + " "
        return text
    except Exception as e:
        print("PDF extraction error:", e)
        return ""

def extract_text_from_image(file_bytes: bytes) -> str:
    try:
        image = vision.Image(content=file_bytes)
        response = vision_client.text_detection(image=image)
        annotations = response.text_annotations

        if annotations:
            return annotations[0].description;
        else:
            return "";
    except Exception as e:
        print("IMAGE extraction error:", e)
        return ""

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX file using python-docx."""
    try:
        document = docx.Document(io.BytesIO(file_bytes))
        full_text = []
        for para in document.paragraphs:
            full_text.append(para.text)
        return "\n".join(full_text)
    except Exception as e:
        print("DOCX extraction error:", e)
        return ""


def extract_text_from_xlsx(file_bytes: bytes) -> str:
    """Extract text from an XLSX file using openpyxl."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        text = ""
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " ".join([str(cell) for cell in row if cell is not None])
                text += row_text + " "
        return text
    except Exception as e:
        print("XLSX extraction error:", e)
        return ""


IMAGE_EXTENSIONS = {'.jpg', '.jpeg', '.png', '.webp', '.gif', '.bmp', '.tiff'}
EXCEL_EXTENSIONS = {'.xlsx', '.xls'}
CSV_EXTENSIONS = {'.csv'}
TABULAR_EXTENSIONS = EXCEL_EXTENSIONS | CSV_EXTENSIONS

MIME_TYPE_MAP = {
    '.pdf': 'application/pdf',
    '.jpg': 'image/jpeg',
    '.jpeg': 'image/jpeg',
    '.png': 'image/png',
    '.webp': 'image/webp',
    '.gif': 'image/gif',
    '.bmp': 'image/bmp',
    '.tiff': 'image/tiff',
    '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    '.xls': 'application/vnd.ms-excel',
    '.csv': 'text/csv',
}


def get_file_category(filename: str) -> str:
    """Return 'pdf', 'image', 'spreadsheet', or 'unknown' based on file extension."""
    ext = os.path.splitext(filename)[1].lower() if filename else ""
    if ext == '.pdf':
        return 'pdf'
    if ext in IMAGE_EXTENSIONS:
        return 'image'
    if ext in TABULAR_EXTENSIONS:
        return 'spreadsheet'
    return 'unknown'


def get_mime_type(filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower() if filename else ""
    return MIME_TYPE_MAP.get(ext, 'application/octet-stream')